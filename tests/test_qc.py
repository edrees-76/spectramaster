"""Detector QC: rules, decay correction, split-peak recombination, persistence, import, GUI/PDF."""
from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import math
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pytest

from spectramaster.model import Spectrum
from spectramaster.tools import qc

T0 = datetime(2024, 1, 1, 9, 0, 0)
CAL = [0.0, 0.5]          # keV = 0.5 * channel
BASE = [100.0 + (1 if i % 2 == 0 else -1) for i in range(20)]   # mean 100, SD 1.026


# ------------------------------------------------------------------ helpers
def _meas(i: int, rate: float = 100.0, ch: float = 1323.314, fw_ch: float = 3.0, live: float = 1000.0,
          coeffs=None, extra_peaks=()) -> qc.QCMeasurement:
    coeffs = list(coeffs or CAL)
    e = sum(c * ch ** k for k, c in enumerate(coeffs))
    peaks = [{"ch": ch, "e": e, "fw_ch": fw_ch, "fw_kev": fw_ch * coeffs[1], "net": rate * live,
              "unc": math.sqrt(rate * live), "sig": 50.0}, *extra_peaks]
    return qc.QCMeasurement(id=f"m{i:03d}", date=T0 + timedelta(days=i), source_path=f"q{i}.spe",
                            live_time=live, real_time=live * 1.01, energy_coeffs=coeffs,
                            total_counts=1e6, window_counts=500.0 * live, peaks=peaks)


def _track(rates, **kw) -> qc.QCTrack:
    t = qc.QCTrack(name="test", detector="D1", source="Cs-137",
                   lines=[qc.MonitoredLine(661.657, "")],    # no nuclide -> no decay correction
                   rate_tol_pct=None, centroid_tol_kev=None, **kw)
    t.measurements = [_meas(i, r) for i, r in enumerate(rates)]
    return t


def _rate_flags(t: qc.QCTrack):
    ev = qc.evaluate(t)
    s = ev.get("rate", 661.657)
    return ev, s, [p.flags for p in s.points]


# ------------------------------------------------------------------ statistics & rules
def test_baseline_statistics_and_clean_series():
    ev, s, flags = _rate_flags(_track(BASE))
    assert s.n_baseline == 20
    assert s.mean == pytest.approx(100.0)
    assert s.sd == pytest.approx(math.sqrt(20 / 19))
    lo, hi = s.limits(3)
    assert hi - lo == pytest.approx(6 * s.sd)
    assert all(f == [] for f in flags)
    assert set(ev.status.values()) == {qc.STATUS_OK}


@pytest.mark.parametrize("tail, rule, where", [
    ([104.0], "1-3s", -1),
    ([102.3, 102.3], "2-2s", -1),
    ([102.3, 97.7], "R-4s", -1),
    ([101.3] * 4, "4-1s", -1),
    ([100.5] * 10, "10-x", -1),
])
def test_westgard_rules_fire(tail, rule, where):
    ev, s, flags = _rate_flags(_track(BASE + tail))
    assert rule in flags[where]
    assert s.points[where].status == qc.STATUS_FAIL
    assert ev.status[s.points[where].mid] == qc.STATUS_FAIL
    # baseline points stay clean
    assert all(f == [] for f in flags[:20])


def test_rule_precision():
    # one point beyond 2 sigma is only a warning; 4-1s needs four points; 10-x needs ten
    _, s, flags = _rate_flags(_track(BASE + [102.3, 101.3, 101.3, 101.3]))
    assert flags[20] == ["1-2s"] and s.points[20].status == qc.STATUS_WARN
    assert "4-1s" in flags[23] and "4-1s" not in flags[22]
    _, s, flags = _rate_flags(_track(BASE + [100.5] * 9))
    assert all("10-x" not in f for f in flags)
    # a disabled rule does not fire
    t = _track(BASE + [104.0])
    t.rules = ["2-2s"]
    _, s, flags = _rate_flags(t)
    assert "1-3s" not in flags[-1] and s.points[-1].status == qc.STATUS_WARN


def test_excluded_points_ignored_and_baseline_period():
    t = _track(BASE + [104.0])
    t.measurements[-1].excluded = True
    ev, s, flags = _rate_flags(t)
    assert flags[-1] == [] and ev.status[t.measurements[-1].id] == qc.STATUS_EXCLUDED
    # an outlier excluded from the baseline does not inflate the SD
    rates = BASE[:10] + [150.0] + BASE[10:]
    t = _track(rates)
    t.measurements[10].excluded = True
    _, s, _ = _rate_flags(t)
    assert s.sd == pytest.approx(math.sqrt(20 / 19)) and s.n_baseline == 20
    # baseline defined by a period
    t = _track(BASE + [110.0] * 5)
    t.baseline_start, t.baseline_end = T0, T0 + timedelta(days=9, hours=1)
    _, s, _ = _rate_flags(t)
    assert s.n_baseline == 10 and s.mean == pytest.approx(100.0)


def test_tolerances_and_violation_runs():
    t = _track(BASE + [101.5, 101.5, 101.5])
    t.rate_tol_pct = 1.0
    t.rules = []
    ev, s, flags = _rate_flags(t)
    assert all("TOL" in f for f in flags[20:])
    assert s.tol_lo == pytest.approx(99.0) and s.tol_hi == pytest.approx(101.0)
    runs = [r for r in qc.violation_runs(ev) if r.rule == "TOL"]
    assert len(runs) == 1 and runs[0].n == 3 and runs[0].first == T0 + timedelta(days=20)
    # FWHM limit
    t = _track(BASE)
    t.lines[0].fwhm_max = 1.4                # fwhm = 3 ch * 0.5 keV = 1.5 keV
    ev = qc.evaluate(t)
    assert all("TOL" in p.flags for p in ev.get("fwhm", 661.657).points)


def test_gain_drift_detected_only_with_fixed_calibration():
    """Each day the lab re-calibrated (file energy stays 661.66), but the peak walks in
    channels: only the fixed reference calibration exposes the drift."""
    t = qc.QCTrack(name="drift", lines=[qc.MonitoredLine(661.657, "")], centroid_tol_kev=0.5,
                   rate_tol_pct=None)
    for i in range(30):
        ch = 1323.314 + (0.0 if i < 20 else 0.4 * (i - 19))          # drift after day 20
        gain = 661.657 / ch
        t.measurements.append(_meas(i, BASE[i % 20], ch=ch, coeffs=[0.0, gain]))
    t.centroid_mode = "file"
    s = qc.evaluate(t).get("centroid", 661.657)
    assert all(p.value == pytest.approx(661.657) for p in s.points)
    assert not any("TOL" in p.flags for p in s.points)
    t.centroid_mode = "fixed"                  # reference calibration = first measurement
    s = qc.evaluate(t).get("centroid", 661.657)
    assert s.points[-1].value == pytest.approx(0.5 * (1323.314 + 4.0), abs=1e-6)
    tol = [i for i, p in enumerate(s.points) if "TOL" in p.flags]
    assert tol and min(tol) == 22              # 0.5 keV = 1 channel -> exceeded from day 22 (1.2 ch)
    assert s.slope_per_year() > 0


def test_split_peak_recombined_and_flagged():
    comp = {"ch": 1328.0, "e": 664.0, "fw_ch": 3.0, "fw_kev": 1.5, "net": 50_000.0, "unc": 230.0, "sig": 40}
    t = _track(BASE)
    m = _meas(20, 50.0, ch=1322.0, extra_peaks=[comp])     # two halves 6 ch apart
    t.measurements.append(m)
    t.match_tol_kev = 3.0
    ev = qc.evaluate(t)
    r = ev.line_results[m.id]["661.657"]
    assert r.components == 2
    assert r.rate == pytest.approx(100.0)                   # areas summed
    assert r.centroid_ch == pytest.approx(1325.0)
    sigma = math.sqrt((3.0 / 2.35482) ** 2 + 9.0)
    assert r.fwhm_kev == pytest.approx(0.5 * 2.35482 * sigma, rel=1e-4)
    assert "SPLIT" in ev.get("fwhm", 661.657).points[-1].flags
    assert ev.status[m.id] == qc.STATUS_FAIL


# ------------------------------------------------------------------ decay correction
def test_decay_factor():
    hl = qc.half_life_of("Cs-137")
    assert hl == pytest.approx(30.08 * 365.25 * 86400, rel=2e-3)
    ref = datetime(2000, 1, 1)
    t = ref + timedelta(seconds=hl)
    assert qc.decay_factor(hl, t, ref) == pytest.approx(2.0)
    assert qc.decay_factor(hl, ref, ref) == 1.0
    assert qc.decay_factor(None, t, ref) == 1.0
    # decay during the count (short-lived): Tr = one half-life -> ln2 / (1 - 1/2)
    assert qc.decay_factor(3600.0, ref, ref, real_time=3600.0) == pytest.approx(2 * math.log(2))


def test_decay_corrected_rate_is_flat_and_efficiency():
    hl = qc.half_life_of("Co-60")
    t = qc.QCTrack(name="co", lines=[qc.MonitoredLine(1332.492, "Co-60")], rate_tol_pct=2.0,
                   reference_date=T0, reference_activity={"Co-60": 10_000.0})
    lam = math.log(2) / hl
    for i in range(25):
        d = T0 + timedelta(days=60 * i)
        rate = 100.0 * math.exp(-lam * (d - T0).total_seconds())
        m = _meas(i, rate, ch=1332.492 / 0.5, live=1000.0)
        m.date, m.real_time = d, 0.0
        t.measurements.append(m)
    ev = qc.evaluate(t)
    s = ev.get("rate", 1332.492)
    assert s.points[-1].value == pytest.approx(100.0, rel=1e-6)      # 4 years later, corrected
    assert not any(p.flags for p in s.points)
    r = ev.line_results[t.measurements[-1].id]["1332.492"]
    assert r.decay == pytest.approx(math.exp(lam * 1440 * 86400))
    assert r.efficiency == pytest.approx(100.0 / (10_000.0 * 0.9998), rel=1e-3)
    # uncorrected the same series would drift out of tolerance
    t.lines[0].half_life_s = 1e30
    s = qc.evaluate(t).get("rate", 1332.492)
    assert "TOL" in s.points[-1].flags


# ------------------------------------------------------------------ persistence
def test_store_roundtrip(tmp_path):
    store = qc.QCStore(tmp_path / "qc")
    t = _track(BASE + [104.0], reference_date=T0, reference_activity={"Cs-137": 3.7e4})
    t.set_lines_from_preset("Cs-137 + Co-60")
    t.baseline_start = T0
    t.measurements[3].excluded = True
    t.measurements[3].note = "source moved — تجربة"
    t.bkg_window = (50.0, None)
    p = store.save(t)
    assert p.exists() and p.parent == tmp_path / "qc"
    t2 = store.load(t.id)
    assert t2.to_json() == t.to_json()
    assert t2.measurements[3].note.endswith("تجربة") and t2.measurements[3].excluded
    assert [l.energy for l in t2.lines] == [661.657, 1173.228, 1332.492]
    assert t2.lines[-1].fwhm_max == 2.0
    e1, e2 = qc.evaluate(t), qc.evaluate(t2)
    assert e1.status == e2.status and e1.flags == e2.flags
    assert [x.name for x in store.list_tracks()] == ["test"]
    store.delete(t2)
    assert store.list_tracks() == []


def test_default_dir_under_appdata(monkeypatch, tmp_path):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    assert qc.default_qc_dir() == tmp_path / "SpectraMaster" / "qc"


# ------------------------------------------------------------------ import of real spectrum files
def _synthetic(day: int, area: float = 40_000.0, shift_ch: float = 0.0, seed: int = 0) -> Spectrum:
    rng = np.random.default_rng(seed)
    n = 4096
    ch = np.arange(n, dtype=float)
    y = 30.0 + 2000.0 / (1 + ch / 300.0)
    for e, a in ((661.657, area), (1173.228, area / 6), (1332.492, area / 7)):
        c = e / 0.5 + shift_ch
        sig = (math.sqrt(0.75 + 0.00185 * e) / 0.5) / 2.3548
        y += a / (sig * math.sqrt(2 * math.pi)) * np.exp(-0.5 * ((ch - c) / sig) ** 2)
    counts = rng.poisson(y).astype(float).tolist()
    return Spectrum(counts=counts, energy_coeffs=[0.0, 0.5], live_time=1000.0, real_time=1005.0,
                    start_time=T0 + timedelta(days=day), sample_id=f"QC day {day}", detector="HPGe-T")


def _write_series(folder: Path) -> list[Path]:
    from spectramaster.formats import get_format, load_all
    load_all()
    w = get_format("iaea_spe").writer
    folder.mkdir(parents=True, exist_ok=True)
    out = []
    for k, day in enumerate([3, 0, 4, 1, 2]):          # file names not in date order
        area = 20_000.0 if day == 4 else 40_000.0      # injected outlier: half the activity
        p = folder / f"qc_{k}.spe"
        w(_synthetic(day, area, seed=day), p)
        out.append(p)
    return out


def test_import_folder_sorted_dedup_and_outlier(tmp_path):
    files = _write_series(tmp_path / "series")
    (tmp_path / "series" / "copy.spe").write_bytes(files[0].read_bytes())     # duplicate
    (tmp_path / "series" / "junk.spe").write_text("not a spectrum")
    t = qc.QCTrack(name="imp", detector="HPGe-T", match_tol_kev=1.5, baseline_n=4)
    t.set_lines_from_preset("Cs-137 + Co-60")
    calls = []
    res = qc.import_paths(t, [tmp_path / "series"], progress=lambda i, n, p: calls.append((i, n)))
    assert len(res.added) == 5 and len(res.duplicates) == 1 and len(res.errors) == 1
    assert calls[-1] == (7, 7)
    dates = [m.date for m in t.measurements]
    assert dates == sorted(dates) and dates[0] == T0
    ev = qc.evaluate(t)
    rates = [p.value for p in ev.get("rate", 661.657).points]
    assert all(r is not None for r in rates)
    assert rates[0] == pytest.approx(40.0, rel=0.05)      # 40000 counts / 1000 s (Cs-137 decay ~0)
    assert rates[4] == pytest.approx(20.0, rel=0.05)
    assert ev.status[t.measurements[4].id] == qc.STATUS_FAIL          # TOL (±5 %) and 1-3s
    assert "rate@661.657:TOL" in ev.flags[t.measurements[4].id]
    fw = ev.get("fwhm", 1332.492).points[0].value
    assert fw == pytest.approx(math.sqrt(0.75 + 0.00185 * 1332.492), rel=0.1)
    # re-analysis from the files keeps notes / exclusions
    t.measurements[1].note = "kept"
    r2 = qc.reanalyse(t)
    assert len(r2.added) == 5 and t.measurements[1].note == "kept"


def test_add_spectrum_duplicate_and_xlsx(tmp_path):
    t = qc.QCTrack(name="x")
    t.set_lines_from_preset("Cs-137")
    s = _synthetic(0)
    assert qc.add_spectrum(t, s) is not None
    assert qc.add_spectrum(t, s) is None
    out = qc.export_track_xlsx(t, tmp_path / "x.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["Track", "Measurements", "Statistics", "Violations"]
    assert wb["Measurements"].max_row == 2


# ------------------------------------------------------------------ GUI
@pytest.fixture(scope="module")
def qapp():
    pytest.importorskip("PySide6")
    from spectramaster.gui.app import create_app
    return create_app(["test"])


def _wait(cond, timeout=60.0):
    import time
    from PySide6.QtWidgets import QApplication
    end = time.monotonic() + timeout
    while not cond() and time.monotonic() < end:
        QApplication.processEvents()
        time.sleep(0.01)
    assert cond()


def test_qc_window_offscreen_and_exports(qapp, tmp_path):
    from spectramaster.gui.qc_charts import chart_image, paint_qc_report
    from spectramaster.gui.qc_window import QCWindow, parse_lines, lines_to_text
    store = qc.QCStore(tmp_path / "store")
    w = QCWindow(store=store)
    try:
        assert w.track is None
        t = qc.QCTrack(name="GUI track", detector="HPGe-T", match_tol_kev=1.5, baseline_n=3)
        t.set_lines_from_preset("Cs-137 + Co-60")
        w.add_track(t)
        assert w.track is not None and w.track.id == t.id
        files = _write_series(tmp_path / "series")
        res = w.add_paths(files, background=False)
        assert len(res.added) == 5
        assert w.table.rowCount() == 5
        assert w.charts["rate"].series is not None and len(w.charts["rate"].series.points) == 5
        assert w.viol_table.rowCount() >= 1
        # persisted
        assert len(store.load(t.id).measurements) == 5
        # add the "displayed" spectrum
        m = w.add_current_spectrum(_synthetic(9, seed=9))
        assert m is not None and w.table.rowCount() == 6
        # exclusion from the table
        w.set_excluded([m.id], True)
        assert w.ev.status[m.id] == qc.STATUS_EXCLUDED
        w.select_measurement(m.id)
        assert w.selected_ids() == [m.id]
        w.line_combo.setCurrentIndex(2)
        assert w.charts["centroid"].series.line.energy == 1332.492
        w.grab()                                           # paints the charts
        pdf = w.export_pdf(tmp_path / "out" / "report.pdf")      # written on a worker thread
        xl = w.export_excel(tmp_path / "out" / "qc.xlsx")
        assert w.exports_running()
        _wait(lambda: not w.exports_running())
        assert pdf.exists() and pdf.read_bytes()[:4] == b"%PDF" and pdf.stat().st_size > 5000
        assert xl.exists()
        pngs = w.export_charts(tmp_path / "png", "png")
        svgs = w.export_charts(tmp_path / "svg", "svg")
        assert len(pngs) == 10 and all(p.stat().st_size > 1000 for p in pngs)
        assert svgs[0].read_text(encoding="utf-8").lstrip().startswith("<?xml")
        img = chart_image(w.ev.get("fwhm", 661.657), 600, 250)
        assert not img.isNull() and img.width() == 600
        # reusable report painter on a QPdfWriter (what the print module uses)
        from PySide6.QtGui import QPdfWriter
        pw = QPdfWriter(str(tmp_path / "p2.pdf"))
        assert paint_qc_report(pw, w.track) >= 2
        assert parse_lines(lines_to_text(t.lines))[2].fwhm_max == 2.0
        with pytest.raises(ValueError):
            parse_lines("abc")
        w.delete_track(confirm=False)
        assert store.list_tracks() == [] and w.track is None
    finally:
        w.close()
        w.deleteLater()


def test_track_dialog_roundtrip(qapp):
    from spectramaster.gui.qc_window import TrackDialog
    t = qc.QCTrack(name="dlg", detector="G4", reference_activity={"Cs-137": 1000.0})
    t.set_lines_from_preset("Co-60")
    d = TrackDialog(t)
    d.source.setCurrentText("Eu-152")
    d._apply_preset()
    d.rule_boxes["10-x"].setChecked(False)
    d.accept()
    assert [l.nuclide for l in t.lines] == ["Eu-152"] * 4
    assert "10-x" not in t.rules and t.reference_activity == {"Cs-137": 1000.0}
    d.deleteLater()


def test_main_window_menu_entry(qapp, tmp_path):
    from PySide6.QtCore import QSettings
    from PySide6.QtWidgets import QMenu
    from spectramaster.gui.app import MainWindow
    from spectramaster.gui.settings import Settings
    s = Settings(QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat))
    s.set("qc/dir", str(tmp_path / "qcdir"))
    mw = MainWindow(s)
    try:
        menu = mw.menuBar().findChild(QMenu, "menu_tools")
        assert menu is not None and mw.act_qc in menu.actions()
        mw.act_qc.trigger()
        w = mw._qc_window
        assert w.store.root == tmp_path / "qcdir"
        mw.apply_language("en", save=False)
        assert "QC" in mw.act_qc.text()
        w.close()
    finally:
        mw.close()
        mw.deleteLater()
