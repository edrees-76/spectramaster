"""Archive search: SQLite index (incremental, duplicates, Arabic FTS), query API and GUI window."""
from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import math
import shutil
import time
from datetime import date, datetime
from pathlib import Path

import pytest

from spectramaster.convert import write_spectrum
from spectramaster.formats import get_format, load_all
from spectramaster.model import Spectrum
from spectramaster.tools import archive_db
from spectramaster.tools.archive_db import ArchiveIndex, canonical_nuclide, export_rows, normalize_text


# --------------------------------------------------------------------------- helpers
def make_spectrum(sid="S", desc="", year=2009, month=3, peak_kev=661.66, area=5000.0, live=1000.0,
                  detector="HPGe-A", calibrated=True, seed=0) -> Spectrum:
    n = 2048
    counts = []
    for i in range(n):
        e = 0.5 + 0.75 * i
        c = 40.0 * math.exp(-e / 600.0) + 5 + (seed % 3)
        if peak_kev:
            sig = 1.8 / 2.355 / 0.75
            c += area / (sig * math.sqrt(2 * math.pi)) * math.exp(-0.5 * ((i - (peak_kev - 0.5) / 0.75) / sig) ** 2)
        counts.append(float(round(c)))
    return Spectrum(counts, [0.5, 0.75] if calibrated else [], live, live * 1.01,
                    datetime(year, month, 4, 10, 0, seed % 60), sid, desc, detector)


def write(path: Path, spec: Spectrum, key: str = "iaea_spe") -> Path:
    load_all()
    path.parent.mkdir(parents=True, exist_ok=True)
    write_spectrum(spec, path, get_format(key))
    return path


@pytest.fixture()
def fake_engine(monkeypatch):
    """Fast stand-in for peak search + identification: nuclide decided by the sample ID."""
    from spectramaster import analysis
    from spectramaster.analysis import NuclideMatch, Peak

    calls = []

    def find_peaks(s, **_kw):
        calls.append(s.sample_id)
        return [Peak(centroid=881.5, energy=661.6, fwhm=2.4, fwhm_kev=1.8, net_area=5000.0, net_area_unc=80.0,
                     gross_area=6000.0, background=1000.0, left=875, right=888, significance=60.0,
                     count_rate=5.0)]

    def identify(s, peaks=None, **_kw):
        sid = s.sample_id.lower()
        out = []
        if "cs" in sid:
            out.append(NuclideMatch("Cs-137", 0.9))
        if "weak" in sid:
            out.append(NuclideMatch("Cs-137", 0.3))
        if "k40" in sid:
            out.append(NuclideMatch("K-40", 0.7))
        return out

    monkeypatch.setattr(analysis, "find_peaks", find_peaks)
    monkeypatch.setattr(analysis, "identify", identify)
    return calls


@pytest.fixture()
def archive(tmp_path):
    """Mixed temp archive: SPE / CHN / N42, Arabic names and texts, a mirrored duplicate,
    a non-spectrum text file and a corrupt .s0."""
    root = tmp_path / "الأرشيف"
    a = write(root / "عينات ٢٠٠٩" / "IM-09-1+زَيتون.spe", make_spectrum("cs olive", "عينة الزّيتونة من درعا", 2009, 3))
    write(root / "عينات ٢٠٠٩" / "IM-09-2+شاي.chn", make_spectrum("k40 tea", "", 2009, 7, peak_kev=1460.8, seed=1),
          "ortec_chn")
    write(root / "2010" / "bg.n42", make_spectrum("background", "خلفية المختبر إختبار", 2010, 1, peak_kev=None,
                                                 seed=2), "n42_2012")
    write(root / "2010" / "weak.spe", make_spectrum("weak cs", "ماء الشرب", 2010, 5, seed=3, live=50))
    write(root / "2010" / "uncal.chn", make_spectrum("cs uncal", "", 2010, 6, calibrated=False, seed=4),
          "ortec_chn")
    mirror = root / "نسخة" / "عينات ٢٠٠٩"
    mirror.mkdir(parents=True)
    shutil.copy2(a, mirror / a.name)                       # byte-identical copy
    (root / "2010" / "notes.txt").write_text("not a spectrum at all\nhello\n", encoding="utf-8")
    (root / "2010" / "broken.s0").write_bytes(b"\x00\x01garbage" * 3)
    return root


@pytest.fixture()
def built(tmp_path, archive, fake_engine):
    ix = ArchiveIndex(tmp_path / "db" / "archive.sqlite")
    res = ix.build([archive], workers=1)
    yield ix, res, archive
    ix.close()


# --------------------------------------------------------------------------- unit
def test_normalize_text_arabic():
    assert normalize_text("أحمد إبراهيم آمنة") == normalize_text("احمد ابراهيم امنه")
    assert normalize_text("الزَّيْتُونَة") == "الزيتونه"
    assert normalize_text("مصطفى") == "مصطفي"
    assert normalize_text("شـــاي") == "شاي"                    # tatweel
    assert normalize_text("عينة ٢٠٠٩ IM-09") == "عينه 2009 im-09"
    assert normalize_text(None) == ""


def test_canonical_nuclide():
    names = ["Cs-137", "Ag-110m", "K-40"]
    assert canonical_nuclide("cs137", names) == "Cs-137"
    assert canonical_nuclide("137Cs", names) == "Cs-137"
    assert canonical_nuclide("CS-137", names) == "Cs-137"
    assert canonical_nuclide("ag 110m", names) == "Ag-110m"
    assert canonical_nuclide("K-40") == "K-40"


# --------------------------------------------------------------------------- build
def test_build_mixed_archive(built):
    ix, res, root = built
    st = ix.stats()
    # 6 spectrum files (incl. duplicate) + 1 corrupt .s0; notes.txt is not a spectrum
    assert st["entries"] == 7
    assert st["errors"] == 1
    assert st["duplicates"] == 1
    assert st["unique"] == 5
    assert st["non_spectra"] == 1
    assert res.analysed == 5 and res.analysis_failed == 0
    assert st["pending_analysis"] == 0
    rows = ix.search(include_duplicates=True, limit=None)
    assert len(rows) == 6
    dup = [r for r in rows if r.dup_of is not None]
    assert len(dup) == 1 and dup[0].copies == 2
    canon = next(r for r in rows if r.id == dup[0].dup_of)
    assert canon.sha256 == dup[0].sha256 and canon.path != dup[0].path
    # duplicates analysed once and linked: the duplicate row also shows the nuclides
    assert dup[0].nuclides == canon.nuclides == [("Cs-137", 0.9)]
    olive = next(r for r in rows if r.sample_id == "cs olive")
    assert olive.start_time == datetime(2009, 3, 4, 10, 0, 0)
    assert olive.live_time == pytest.approx(1000.0)
    assert olive.calibrated and olive.energy_coeffs == pytest.approx([0.5, 0.75])
    assert olive.folder == "عينات ٢٠٠٩" and olive.format_key == "iaea_spe"
    assert ix.peaks_of(olive)[0]["energy"] == pytest.approx(661.6)
    assert set(ix.copies_of(olive)) == {canon.path, dup[0].path}


def test_db_inside_archive_refused(archive):
    with pytest.raises(ValueError):
        ArchiveIndex(archive / "index.sqlite").build([archive], workers=1)


def test_incremental_refresh(built, fake_engine):
    ix, _res, root = built
    fake_engine.clear()
    r = ix.build(workers=1)                                  # nothing changed
    assert r.scanned == 0 and r.analysed == 0 and r.unchanged == r.files_listed == 8
    assert fake_engine == []
    # touched (new mtime, same bytes) -> rehash only
    p = root / "2010" / "weak.spe"
    os.utime(p, (time.time() + 10, time.time() + 10))
    # changed content
    write(root / "2010" / "bg.n42", make_spectrum("background k40", "خلفية", 2011, 2, peak_kev=1460.8, seed=5),
          "n42_2012")
    # new file and a removed file
    write(root / "2011" / "new cs.spe", make_spectrum("cs new", "", 2011, 9, seed=6))
    (root / "2010" / "uncal.chn").unlink()
    r = ix.build(workers=1)
    assert r.scanned == 3 and r.rehashed_same == 1 and r.removed == 1
    assert r.analysed == 2                                   # changed + new only
    assert sorted(fake_engine) == ["background k40", "cs new"]
    assert ix.count(year=2011) == 2
    assert [x.sample_id for x in ix.search(nuclides=["K-40"], year=2011)] == ["background k40"]
    assert ix.search(text="uncal") == []
    # removing the duplicate's original promotes the mirror copy to canonical
    (root / "عينات ٢٠٠٩" / "IM-09-1+زَيتون.spe").unlink()
    ix.build(workers=1)
    rows = ix.search(text="زيتون")
    assert len(rows) == 1 and "نسخة" in rows[0].path and rows[0].nuclides == [("Cs-137", 0.9)]


def test_resume_after_cancel(tmp_path, archive, fake_engine):
    ix = ArchiveIndex(tmp_path / "a.sqlite")
    seen = []

    def prog(phase, done, total, eta, cur):
        if phase == "analyse":
            seen.append(done)
            if done >= 2:
                return False
    r = ix.build([archive], workers=1, progress=prog, analysis_batch=1)
    assert r.cancelled and r.analysed == 2 and r.analysis_pending == 3
    r = ix.build(workers=1)
    assert not r.cancelled and r.analysed == 3 and r.analysis_pending == 0
    ix.close()


# --------------------------------------------------------------------------- queries
def test_arabic_text_search(built):
    ix, _res, _root = built
    ids = lambda **kw: sorted(r.sample_id for r in ix.search(**kw))
    assert ids(text="زيتون") == ["cs olive"]                 # file name with tashkeel
    assert ids(text="الزيتونه") == ["cs olive"]              # ة/ه and attached article
    assert ids(text="زيتونة درعا") == ["cs olive"]           # all words
    assert ids(text="اختبار") == ["background"]              # stored as إختبار
    assert ids(text="شاي") == ["k40 tea"]                    # file name
    assert ids(text="٢٠٠٩") == ["cs olive", "k40 tea"]       # folder, Arabic-Indic digits
    assert ids(text="im") == ["cs olive", "k40 tea"]         # short term -> substring fallback
    assert ids(text="nothing-like-this") == []
    assert len(ix.search(text="زيتون", include_duplicates=True)) == 2


def test_nuclide_and_other_filters(built):
    ix, _res, _root = built
    ids = lambda **kw: sorted(r.sample_id for r in ix.search(**kw))
    assert ids(nuclides=["Cs-137"]) == ["cs olive"]          # uncal: no identification, weak: 0.3
    assert ids(nuclides=["cs137"], min_confidence=0.25) == ["cs olive", "weak cs"]
    assert ids(nuclides=["Cs-137", "K-40"], nuclide_mode="any") == ["cs olive", "k40 tea"]
    assert ids(nuclides=["Cs-137", "K-40"]) == []
    assert ids(date_from="2010", date_to="2010-05") == ["background", "weak cs"]
    assert ids(date_from=date(2009, 7, 4), date_to=date(2009, 7, 4)) == ["k40 tea"]
    assert ids(min_live_time=100, year=2010) == ["background", "cs uncal"]
    assert ids(has_calibration=False) == ["cs uncal"]
    assert ids(folder_contains="عينات") == ["cs olive", "k40 tea"]
    assert ids(detector="hpge-a") == ["background", "cs uncal", "k40 tea"]   # CHN / N42 keep the name
    assert len(ids(detector="HPGe")) == 5                              # guessed detector type
    assert len(ids(energy=661.6, energy_tol=0.5)) == 5
    assert ids(energy=1000.0) == []
    rows = ix.search(order_by="-date")
    assert [r.start_time for r in rows] == sorted((r.start_time for r in rows), reverse=True)
    assert ix.search(include_errors=True, has_calibration=None, limit=None, text="broken")[0].status.startswith("error")


def test_aggregates(built):
    ix, _res, _root = built
    assert ix.count_by_period("year") == [("2009", 2), ("2010", 3)]
    assert ix.count_by_period("month", year=2009) == [("2009-03", 1), ("2009-07", 1)]
    assert ix.count(nuclides=["Cs-137"], year=2009) == 1     # "how many samples in 2009 had Cs-137"
    assert dict(ix.nuclide_counts(min_confidence=0.25)) == {"Cs-137": 2, "K-40": 1}
    assert dict(ix.nuclide_counts(year=2010, min_confidence=0.25)) == {"Cs-137": 1}
    assert dict(ix.nuclide_counts(include_duplicates=True)) == {"Cs-137": 2, "K-40": 1}


def test_export_rows(built, tmp_path):
    ix, _res, _root = built
    rows = ix.search()
    csv_path = export_rows(rows, tmp_path / "out" / "r.csv")
    text = csv_path.read_text(encoding="utf-8-sig")
    assert "IM-09-1+زَيتون.spe" in text and "Cs-137 (0.90)" in text
    xlsx = export_rows(rows, tmp_path / "out" / "r.xlsx")
    from openpyxl import load_workbook
    ws = load_workbook(xlsx).active
    assert ws.max_row == len(rows) + 1


def test_process_pool_real_engine(tmp_path):
    """Real pool (spawned workers) + real identification on a small synthetic archive."""
    root = tmp_path / "arc"
    write(root / "a.spe", make_spectrum("A", "", 2009, 1, area=20000))
    write(root / "b.spe", make_spectrum("B", "", 2009, 2, area=20000, seed=1))
    shutil.copy2(root / "a.spe", root / "a copy.spe")
    ix = ArchiveIndex(tmp_path / "p.sqlite")
    phases = set()
    r = ix.build([root], workers=2, progress=lambda ph, d, t, e, c: phases.add(ph))
    assert {"scan", "analyse"} <= phases
    assert r.new_entries == 3 and r.analysed == 2 and r.duplicates == 1
    assert [x.sample_id for x in ix.search(nuclides=["Cs-137"], min_confidence=0.3)] == ["A", "B"]
    ix.close()


# --------------------------------------------------------------------------- GUI
pytest.importorskip("PySide6")


@pytest.fixture(scope="module")
def qapp():
    from spectramaster.gui.app import create_app
    return create_app(["test"])


@pytest.fixture()
def gui_settings(tmp_path):
    from PySide6.QtCore import QSettings
    from spectramaster.gui.settings import Settings
    return Settings(QSettings(str(tmp_path / "settings.ini"), QSettings.IniFormat))


def _wait(cond, timeout=30.0):
    from PySide6.QtWidgets import QApplication
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        QApplication.processEvents()
        if cond():
            return True
        time.sleep(0.01)
    return cond()


def test_gui_window_builds_populates_and_opens(qapp, gui_settings, tmp_path, archive, fake_engine):
    from PySide6.QtCore import Qt
    from spectramaster.gui import i18n
    from spectramaster.gui.app import MainWindow
    from spectramaster.gui.archive_search import ArchiveSearchWindow, show_archive_search

    db = tmp_path / "gui.sqlite"
    gui_settings.set("archive/db", str(db))
    main = MainWindow(gui_settings)
    main.show_error = lambda *_a: None
    try:
        assert main.act_archive_search in main.menu_file.actions()
        assert main.act_archive_search.text() == i18n.STRINGS[i18n.language()]["act_archive_search"]
        win = show_archive_search(main)
        assert isinstance(win, ArchiveSearchWindow) and show_archive_search(main) is win
        win._message = win.messages.append
        assert win.db_path == db
        assert win.add_root(str(archive))
        assert win.start_build(workers=1)
        assert _wait(lambda: not win.is_building())
        assert win.last_build is not None and win.last_build.analysed == 5
        assert win.model.rowCount() == 5                     # duplicates hidden, error row hidden
        # filters
        win.text_edit.setText("الزيتونة")
        rows = win.run_search()
        assert [r.sample_id for r in rows] == ["cs olive"]
        assert "Cs-137" in win.model.index(0, 3).data()
        win.text_edit.setText("")
        win.add_nuclide("cs137")
        win.conf_slider.setValue(25)
        assert sorted(r.sample_id for r in win.run_search()) == ["cs olive", "weak cs"]
        win.chk_dups.setChecked(True)
        assert win.model.rowCount() == 3
        assert win.histogram.data == [("2009", 2), ("2010", 1)]
        # sorting by live time
        win.model.sort(4, Qt.AscendingOrder)
        assert win.model.rows[0].sample_id == "weak cs"
        # double-click opens the spectrum (in place) in the main window
        ix = win.model.index(0, 0)
        win.table.doubleClicked.emit(ix)
        assert main.current_spectrum() is not None
        assert main.current_spectrum().sample_id == "weak cs"
        assert Path(main._path) == Path(win.model.rows[0].path)
        # overlay + copy + export of the selection
        from PySide6.QtCore import QItemSelectionModel
        win.table.selectionModel().select(win.model.index(1, 0),
                                          QItemSelectionModel.ClearAndSelect | QItemSelectionModel.Rows)
        assert [r.path for r in win.selected_rows()] == [win.model.rows[1].path]
        assert win.add_selected_as_overlay() == 1
        assert win.copy_paths() == win.model.rows[1].path
        out = win.export_results(str(tmp_path / "res.csv"))
        assert out.exists() and len(out.read_text(encoding="utf-8-sig").splitlines()) == 4
        # filters + roots remembered
        win.close()
        win2 = ArchiveSearchWindow(main, gui_settings)
        assert win2.roots() == [str(archive)]
        assert win2.selected_nuclides() == ["Cs-137"] and win2.conf_slider.value() == 25
        assert win2.chk_dups.isChecked()
        assert len(win2.run_search()) == 3
        win2.close()
        # language switch retranslates the open window
        main.apply_language("en", save=False)
        assert win.windowTitle() == "Archive Search"
        main.apply_language("ar", save=False)
        assert win.windowTitle() == i18n.STRINGS["ar"]["as_title"]
    finally:
        main.close()
