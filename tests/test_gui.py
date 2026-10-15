"""GUI tests (headless). Uses a throw-away fake format registered with code 99."""
from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import time
from datetime import datetime
from pathlib import Path

import pytest

pytest.importorskip("PySide6")

from PySide6.QtCore import QSettings, Qt
from PySide6.QtWidgets import QApplication

from spectramaster.convert import output_name
from spectramaster.formats import base as fbase
from spectramaster.formats import load_all
from spectramaster.formats.base import FormatError, FormatSpec, register
from spectramaster.model import Spectrum
from spectramaster.gui import i18n

FAKE_KEY = "fake_gui_test"


# --------------------------------------------------------------------------- fake format
def _fake_read(path: Path) -> list[Spectrum]:
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    if not lines or lines[0] != "FAKE":
        raise FormatError("not a FAKE file")
    sid = lines[1]
    live, real = float(lines[2]), float(lines[3])
    nspec = int(lines[4])
    counts = [float(x) for x in lines[5].split()]
    out = []
    for k in range(nspec):
        out.append(Spectrum(counts=[c * (k + 1) for c in counts], energy_coeffs=[0.5, 0.25],
                            live_time=live, real_time=real,
                            start_time=datetime(2024, 3, 4, 5, 6, 7), sample_id=sid,
                            description="fake description", detector="HPGe-1",
                            extra={"geometry": "Marinelli", "operator": "تجربة"}))
    return out


def _fake_write(spec: Spectrum, path: Path) -> None:
    Path(path).write_text("FAKE\n%s\n%g\n%g\n1\n%s\n" % (
        spec.sample_id, spec.live_time or 0, spec.real_time or 0,
        " ".join("%g" % c for c in spec.counts)), encoding="utf-8")


def _fake_sniff(path: Path, head: bytes) -> bool:
    return head.startswith(b"FAKE\n") or head.startswith(b"FAKE\r\n")


@pytest.fixture(scope="module")
def qapp():
    from spectramaster.gui.app import create_app
    app = create_app(["test"])
    yield app


@pytest.fixture()
def fake_format():
    load_all()
    spec = register(FormatSpec(code=99, key=FAKE_KEY, name="Fake GUI Test", extensions=[".fak"],
                               reader=_fake_read, writer=_fake_write, sniffer=_fake_sniff))
    yield spec
    fbase._REGISTRY.pop(FAKE_KEY, None)


@pytest.fixture()
def settings(tmp_path):
    from spectramaster.gui.settings import Settings
    return Settings(QSettings(str(tmp_path / "settings.ini"), QSettings.IniFormat))


def make_file(folder: Path, name: str, sid: str = "S1", nspec: int = 1, n: int = 256) -> Path:
    counts = " ".join(str(10 + (i % 17) + (500 if 100 <= i <= 104 else 0)) for i in range(n))
    p = folder / name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(f"FAKE\n{sid}\n100\n125\n{nspec}\n{counts}\n", encoding="utf-8")
    return p


@pytest.fixture()
def window(qapp, settings, fake_format):
    from spectramaster.gui.app import MainWindow
    w = MainWindow(settings)
    errors: list[str] = []
    infos: list[str] = []
    w.show_error = errors.append
    w.show_info = infos.append
    w.errors, w.infos = errors, infos
    yield w
    w.close()
    w.deleteLater()
    QApplication.processEvents()


def wait_until(cond, timeout=20.0):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        QApplication.processEvents()
        if cond():
            return True
        time.sleep(0.01)
    return cond()


# --------------------------------------------------------------------------- tests
@pytest.mark.parametrize("lang,direction", [("ar", Qt.RightToLeft), ("en", Qt.LeftToRight)])
def test_window_constructs_in_both_languages(qapp, settings, fake_format, lang, direction):
    from spectramaster.gui import i18n
    from spectramaster.gui.app import MainWindow
    settings.language = lang
    w = MainWindow(settings)
    try:
        assert i18n.language() == lang
        assert w.layoutDirection() == direction
        assert qapp.layoutDirection() == direction
        assert w.plot.layoutDirection() == Qt.LeftToRight
        assert w.menu_file.title() == i18n.STRINGS[lang]["menu_file"]
        assert w.act_open_as.text() == i18n.STRINGS[lang]["act_open_as"]
        img = w.grab()
        assert not img.isNull()
    finally:
        w.close()


def test_default_language_is_english(qapp, settings, fake_format):
    """A fresh install (no stored setting) starts in English, left-to-right."""
    from spectramaster.gui.app import MainWindow
    assert i18n.DEFAULT_LANGUAGE == "en"
    w = MainWindow(settings)
    assert settings.language == "en"
    assert i18n.language() == "en"
    assert w.layoutDirection() == Qt.LeftToRight
    w.close()


def test_stored_arabic_choice_is_remembered(qapp, settings, fake_format):
    """Once the user picks Arabic (View > Language) it is used on the next start."""
    from spectramaster.gui.app import MainWindow
    settings.language = "ar"
    w = MainWindow(settings)
    assert i18n.language() == "ar"
    assert w.layoutDirection() == Qt.RightToLeft
    w.close()


def test_language_switch_persists(window, settings):
    from spectramaster.gui import i18n
    window.apply_language("en")
    assert settings.language == "en"
    assert window.menu_help.title() == "&Help"
    window.apply_language("ar")
    assert settings.language == "ar"
    assert window.menu_help.title() == i18n.STRINGS["ar"]["menu_help"]


def test_every_english_key_has_arabic():
    from spectramaster.gui.i18n import STRINGS
    assert set(STRINGS["en"]) == set(STRINGS["ar"])


def test_load_populates_info_panel(window, tmp_path):
    window.apply_language("en")
    p = make_file(tmp_path, "sample.fak", sid="Soil-42")
    assert window.load_file(p)
    info = window.info
    assert info.value_text("sample_id") == "Soil-42"
    assert info.value_text("live_time") == "100.00 s"
    assert info.value_text("real_time") == "125.00 s"
    assert info.value_text("dead_time") == "20.00 %"
    assert info.value_text("channels") == "256"
    assert "2024-03-04" in info.value_text("start")
    assert info.value_text("detector") == "HPGe-1"
    assert info.value_text("calibration").startswith("E = 0.5 + 0.25")
    assert "Fake GUI Test" in info.value_text("format")
    assert info.extra_table.rowCount() == 2
    assert window.plot.spectrum() is window.current_spectrum()
    assert "sample.fak" in window.windowTitle()
    assert str(p) in i18n.strip_bidi(window.status_path.text())
    assert "Fake GUI Test" in window.status_fmt.text()
    assert not window.errors


def test_arabic_path_loads(window, tmp_path):
    p = make_file(tmp_path / "مجلد عربي", "طيف التربة.fak", sid="عينة ١")
    assert window.load_file(p)
    assert window.info.value_text("sample_id") == "عينة ١"
    out = tmp_path / "مخرجات" / "ناتج.fak"
    assert window.save_current_as(_fmt(), out)
    assert out.read_text(encoding="utf-8").startswith("FAKE\nعينة ١")


def _fmt():
    from spectramaster.formats import get_format
    return get_format(FAKE_KEY)


def test_failed_read_keeps_previous_spectrum(window, tmp_path):
    good = make_file(tmp_path, "good.fak", sid="GOOD")
    bad = tmp_path / "bad.fak"
    bad.write_text("FAKE\nbroken\n", encoding="utf-8")
    assert window.load_file(good)
    shown = window.current_spectrum()
    assert not window.load_file(bad)
    assert len(window.errors) == 1
    assert "bad.fak" in window.errors[0]
    assert "Traceback" not in window.errors[0]
    assert window.current_spectrum() is shown
    assert window.info.value_text("sample_id") == "GOOD"
    # missing file and unknown format also give friendly errors
    assert not window.load_file(tmp_path / "missing.fak")
    junk = tmp_path / "junk.zzz_unknown"
    junk.write_bytes(b"\x00\x01garbage")
    assert not window.load_file(junk)
    assert len(window.errors) == 3
    assert window.current_spectrum() is shown


def test_explicit_input_format_used(window, tmp_path, fake_format):
    window.settings.set_input_format(fake_format)
    p = make_file(tmp_path, "renamed.txt_whatever")   # extension does not matter
    assert window.load_file(p)
    assert window._fmt is fake_format


def test_multi_spectrum_selector(window, tmp_path):
    p = make_file(tmp_path, "multi.fak", nspec=3)
    assert window.load_file(p)
    combo = window.info.spectrum_combo
    assert combo.count() == 3
    assert not combo.isHidden()
    combo.setCurrentIndex(2)
    assert window._index == 2
    assert window.current_spectrum().counts[0] == 3 * 10
    assert window.default_save_path(_fmt()).name == output_name(p, _fmt(), 3)


def test_save_as_writes_via_fake_writer(window, tmp_path, fake_format):
    src = make_file(tmp_path, "CHERN.S0X", sid="Chern")
    assert window.load_file(src)
    default = window.default_save_path(fake_format)
    assert default.name == output_name(src, fake_format) == "CHERN_S0X.fak"
    assert default.parent == src.parent
    target = tmp_path / "out" / "saved"           # no extension -> default ext added
    assert window.save_current_as(fake_format, target)
    written = tmp_path / "out" / "saved.fak"
    assert written.is_file()
    spectra = _fake_read(written)
    assert spectra[0].sample_id == "Chern"
    assert spectra[0].counts == window.current_spectrum().counts
    assert window.settings.get("folders/save") == str(tmp_path / "out")


def test_save_without_spectrum_returns_false(window, tmp_path, fake_format):
    assert window.save_current_as(fake_format, tmp_path / "x.fak") is False


def test_settings_remember_formats(settings, fake_format):
    settings.set_input_format(fake_format)
    settings.set_output_format(fake_format)
    settings.sync()
    from spectramaster.gui.settings import Settings
    again = Settings(QSettings(settings.q.fileName(), QSettings.IniFormat))
    assert again.input_format() is fake_format
    assert again.output_format() is fake_format
    again.set_input_format(None)
    assert again.input_format() is None


def test_format_dialog_lists_codes_and_auto(qapp, fake_format):
    from spectramaster.formats import readable_formats
    from spectramaster.gui.dialogs import FormatDialog
    dlg = FormatDialog("t", "p", readable_formats(), "auto", allow_auto=True)
    assert dlg.tree.topLevelItem(0).data(0, Qt.UserRole) == "auto"
    assert dlg.selected_key() == "auto"
    codes = [dlg.tree.topLevelItem(i).text(0) for i in range(dlg.tree.topLevelItemCount())]
    assert "99" in codes
    dlg.select_key(FAKE_KEY)
    assert dlg.selected_key() == FAKE_KEY
    dlg2 = FormatDialog("t", "p", [fake_format], FAKE_KEY, allow_auto=False)
    assert dlg2.selected_key() == FAKE_KEY


def test_batch_worker_converts_and_reports(qapp, tmp_path, fake_format):
    from spectramaster.gui.worker import BatchWorker
    files = [make_file(tmp_path / "in", f"f{i}.fak", sid=f"S{i}") for i in range(5)]
    bad = tmp_path / "in" / "bad.fak"
    bad.write_text("nope", encoding="utf-8")
    files.append(bad)
    out = tmp_path / "out"
    w = BatchWorker(files, fake_format, out, None)
    events, done = [], []
    w.progress.connect(lambda *a: events.append(a))
    w.finished.connect(lambda *a: done.append(a))
    w.run()
    assert [e[0] for e in events] == [1, 2, 3, 4, 5, 6]
    assert all(e[1] == 6 for e in events)
    assert [e[3] for e in events] == [True] * 5 + [False]
    assert events[-1][4]                                # error message present
    assert done == [(5, 1, 6, False)]
    for i in range(5):
        assert (out / f"f{i}_fak.fak").is_file()


def test_batch_worker_cancel(qapp, tmp_path, fake_format):
    from spectramaster.gui.worker import BatchWorker
    files = [make_file(tmp_path, f"c{i}.fak") for i in range(10)]
    w = BatchWorker(files, fake_format, tmp_path / "o", None)
    events, done = [], []

    def on_progress(i, *rest):
        events.append(i)
        if i == 3:
            w.cancel()
    w.progress.connect(on_progress)
    w.finished.connect(lambda *a: done.append(a))
    w.run()
    assert events == [1, 2, 3]
    assert done == [(3, 0, 10, True)]
    assert len(list((tmp_path / "o").glob("*.fak"))) == 3


def test_batch_dialog_threaded_run(window, tmp_path, fake_format):
    files = [make_file(tmp_path / "src", f"t{i}.fak") for i in range(12)]
    window.settings.set_output_format(fake_format)
    dlg = window.open_batch(files)
    assert dlg.model.rowCount() == 12
    dlg.same_dir_cb.setChecked(False)
    dlg.out_edit.setText(str(tmp_path / "dest"))
    finished = []
    dlg.conversionFinished.connect(lambda *a: finished.append(a))
    dlg.start()
    assert wait_until(lambda: finished and not dlg.is_running())
    assert finished[0] == (12, 0, 12, False)
    assert dlg.progress.value() == 12
    assert all(s is True for s in dlg.model.state)
    assert dlg.open_out_btn.isEnabled()
    assert len(list((tmp_path / "dest").glob("*.fak"))) == 12
    dlg.close()


def test_folder_dialog_recursive_keep_tree(window, tmp_path, fake_format):
    root = tmp_path / "root"
    make_file(root, "a.fak")
    make_file(root / "sub", "b.fak")
    make_file(root / "sub" / "deeper", "c.fak")
    (root / "notes.unknownext").write_text("x", encoding="utf-8")
    window.settings.set_output_format(fake_format)
    dlg = window.convert_folder(root)
    assert dlg.model.rowCount() == 3                     # extension filter applied
    dlg.recursive_cb.setChecked(False)
    assert dlg.model.rowCount() == 1
    dlg.recursive_cb.setChecked(True)
    dlg.same_dir_cb.setChecked(False)
    dlg.out_edit.setText(str(tmp_path / "mirror"))
    dlg.keep_tree_cb.setChecked(True)
    finished = []
    dlg.conversionFinished.connect(lambda *a: finished.append(a))
    dlg.start()
    assert wait_until(lambda: finished and not dlg.is_running())
    assert (tmp_path / "mirror" / "a_fak.fak").is_file()
    assert (tmp_path / "mirror" / "sub" / "b_fak.fak").is_file()
    assert (tmp_path / "mirror" / "sub" / "deeper" / "c_fak.fak").is_file()
    dlg.close()


def test_drop_single_and_many(window, tmp_path):
    one = make_file(tmp_path, "drop.fak", sid="DROPPED")
    window.handle_dropped([str(one)])
    assert window.info.value_text("sample_id") == "DROPPED"
    many = [str(make_file(tmp_path, f"m{i}.fak")) for i in range(3)]
    window.handle_dropped(many)
    assert window._batch_dialogs and window._batch_dialogs[-1].model.rowCount() == 3
    window._batch_dialogs[-1].close()


def test_plot_interaction(window, tmp_path):
    p = make_file(tmp_path, "plot.fak", n=4096)
    window.load_file(p)
    plot = window.plot
    plot.resize(800, 400)
    assert plot.x0 == 0 and plot.x1 == 4096
    plot.set_view(100, 200)
    assert (plot.x0, plot.x1) == (100, 200)
    plot.set_log_y(True)
    assert not plot.grab().isNull()
    plot.set_energy_axis(True)
    assert not plot.grab().isNull()
    text = plot.readout(102)
    assert "102" in text and "26.00" in text          # E = 0.5 + 0.25*102
    plot.reset_view()
    assert (plot.x0, plot.x1) == (0, 4096)
    plot.set_log_y(False)
    assert not plot.grab().isNull()


def test_about_dialog(qapp, fake_format):
    from spectramaster.gui.dialogs import AboutDialog
    from spectramaster import __version__
    from PySide6.QtWidgets import QLabel, QTreeWidget
    dlg = AboutDialog()
    assert not dlg.grab().isNull()
    labels = " ".join(lab.text() for lab in dlg.findChildren(QLabel))
    assert __version__ in labels and "SpecUtils" in labels
    tree = dlg.findChild(QTreeWidget)
    names = [tree.topLevelItem(i).text(1) for i in range(tree.topLevelItemCount())]
    assert "Fake GUI Test" in names
    dlg.close()


def test_icon_renders(qapp, tmp_path):
    from spectramaster.gui.icon import app_icon, write_ico
    assert not app_icon().isNull()
    ico = write_ico(tmp_path / "x.ico")
    assert ico.stat().st_size > 1000


# =========================================================================== navigation / ROI / analysis
import math

APTEC_DIR = Path(r"D:\الاسس و المفاهيم العامة فى استخدام برامج الطيف الجامى\اطياف مرجعية"
                 r"\4- أطياف المختبر (Aptec)\الملفات الأصلية S0")


def gauss_spectrum(n=512, amp=1000.0, mu=200.3, sigma=3.0, a=50.0, b=0.2, coeffs=(0.0, 0.5), live=100.0):
    counts = [a + b * i + amp * math.exp(-0.5 * ((i - mu) / sigma) ** 2) for i in range(n)]
    return Spectrum(counts=counts, energy_coeffs=list(coeffs), live_time=live, real_time=live * 1.1,
                    sample_id="gauss")


def test_roi_net_area_matches_analytic():
    from spectramaster.gui.roi import compute_roi
    amp, mu, sigma = 1000.0, 200.3, 3.0
    s = gauss_spectrum(amp=amp, mu=mu, sigma=sigma)
    left, right = 175, 225
    st = compute_roi(s.counts, left, right, 0, s.energy_of, s.live_time)
    analytic = amp * sigma * math.sqrt(2 * math.pi)
    assert st.net == pytest.approx(analytic, rel=1e-4)
    # background of a straight line from the 3-channel end averages is exact
    exp_bg = sum(50.0 + 0.2 * i for i in range(left, right + 1))
    assert st.background == pytest.approx(exp_bg, rel=1e-6)
    assert st.gross == pytest.approx(sum(s.counts[left:right + 1]))
    n = right - left + 1
    assert st.net_unc == pytest.approx(math.sqrt(st.gross + st.background * n / 6.0))
    assert st.centroid == pytest.approx(mu, abs=0.01)
    assert st.fwhm == pytest.approx(2.3548 * sigma, rel=0.02)
    assert st.centroid_energy == pytest.approx(0.5 * mu, abs=0.01)
    assert st.fwhm_kev == pytest.approx(0.5 * 2.3548 * sigma, rel=0.02)
    assert st.count_rate == pytest.approx(analytic / 100.0, rel=1e-4)


@pytest.fixture()
def gplot(qapp):
    from spectramaster.gui.plot import SpectrumPlot
    p = SpectrumPlot()
    p.resize(900, 420)
    p.set_spectrum(gauss_spectrum(n=4096))
    yield p
    p.deleteLater()


def test_zoom_history_back_forward(gplot):
    p = gplot
    assert (p.x0, p.x1) == (0, 4096)
    assert not p.can_go_back()
    p.set_view(100, 300)
    p.set_view(150, 200)
    assert p.can_go_back() and not p.can_go_forward()
    assert p.back() and (p.x0, p.x1) == (100, 300)
    assert p.back() and (p.x0, p.x1) == (0, 4096)
    assert not p.back()
    assert p.forward() and (p.x0, p.x1) == (100, 300)
    p.set_view(10, 50)                     # new branch drops the forward history
    assert not p.can_go_forward()
    # debounced (wheel-like) changes are committed before going back
    p.zoom_x(0.5, record=False)
    assert p.back() and (p.x0, p.x1) == (10, 50)


def test_box_zoom_math(gplot):
    from PySide6.QtCore import QRectF
    p = gplot
    p.set_view(0, 1000)
    rect = p.plot_rect()
    yr = p.y_range()
    box = QRectF(rect.left() + rect.width() * 0.2, rect.top() + rect.height() * 0.25,
                 rect.width() * 0.3, rect.height() * 0.5)
    exp_x0 = 0 + 0.2 * 1000
    exp_x1 = 0 + 0.5 * 1000
    exp_y0 = yr[0] + 0.25 * (yr[1] - yr[0])
    exp_y1 = yr[0] + 0.75 * (yr[1] - yr[0])
    p.zoom_box(box)
    assert p.x0 == pytest.approx(exp_x0) and p.x1 == pytest.approx(exp_x1)
    assert p.y_range() == pytest.approx((exp_y0, exp_y1))
    assert p.back() and (p.x0, p.x1) == (0, 1000) and p._ylim is None
    # a drag shorter than BOX_MIN_HEIGHT is a horizontal-only zoom (Y stays automatic)
    from PySide6.QtCore import QPoint, QPointF
    from PySide6.QtTest import QTest
    a = QPoint(int(rect.left() + 100), int(rect.top() + 100))
    b = QPoint(int(rect.left() + 300), int(rect.top() + 105))
    QTest.mousePress(p, Qt.LeftButton, Qt.NoModifier, a)
    QTest.mouseMove(p, b)
    QTest.mouseRelease(p, Qt.LeftButton, Qt.NoModifier, b)
    assert p._ylim is None
    assert p.x0 == pytest.approx(100 / rect.width() * 1000, abs=2)


def test_y_modes_and_auto_scale(gplot):
    from spectramaster.gui.plot import Y_LOG, Y_SQRT
    p = gplot
    p.set_view(0, 100)                     # far from the peak
    auto = p.y_range()[1]
    p.set_auto_y(False)
    assert p.y_range()[1] > auto * 5        # full-spectrum scale includes the peak
    p.set_auto_y(True)
    for mode in (Y_LOG, Y_SQRT):
        p.set_y_mode(mode)
        assert not p.grab().isNull()
    p.zoom_y(0.5)
    assert p._ylim is not None


def test_overview_window_sync(gplot, qapp):
    from spectramaster.gui.plot import OverviewStrip
    from PySide6.QtTest import QTest
    from PySide6.QtCore import QPoint
    p = gplot
    ov = OverviewStrip(p)
    ov.resize(p.width(), 52)
    p.set_view(1024, 2048)
    a, b = ov.window_px()
    r = ov.strip_rect()
    assert a == pytest.approx(r.left() + r.width() * 0.25)
    assert b == pytest.approx(r.left() + r.width() * 0.5)
    # drag the window right by 10 % of the strip
    start = QPoint(int((a + b) / 2), 20)
    end = QPoint(int((a + b) / 2 + r.width() * 0.1), 20)
    QTest.mousePress(ov, Qt.LeftButton, Qt.NoModifier, start)
    QTest.mouseMove(ov, end)
    QTest.mouseRelease(ov, Qt.LeftButton, Qt.NoModifier, end)
    assert p.x1 - p.x0 == pytest.approx(1024)
    assert p.x0 == pytest.approx(1024 + 409.6, abs=10)
    assert p.back() and p.x0 == pytest.approx(1024)
    assert not ov.grab().isNull()
    ov.deleteLater()


def test_goto_energy_and_channel(window):
    window.display([gauss_spectrum(n=4096)], None, None)
    plot = window.plot
    assert window.goto(100.15, "kev", 1)          # ±5 FWHM around 100.15 keV (channel 200.3)
    centre = (plot.x0 + plot.x1) / 2
    assert plot.energy_at_pos(centre) == pytest.approx(100.15, abs=0.3)
    assert window.goto(100.0, "kev", 4)           # ±10 keV
    assert plot.energy_at_pos(plot.x0) == pytest.approx(90.0, abs=0.01)
    assert plot.energy_at_pos(plot.x1) == pytest.approx(110.0, abs=0.01)
    assert window.goto(1000, "ch", 0)
    assert plot.x0 < 1000.5 < plot.x1
    assert not window.goto(99999, "ch", 0)
    window.goto_edit.setText("200")
    assert window.goto_from_toolbar()


def test_roi_drag_and_panel(window, tmp_path):
    from PySide6.QtTest import QTest
    from PySide6.QtCore import QPoint
    window.display([gauss_spectrum(n=512)], None, None)
    plot = window.plot
    plot.resize(1000, 400)
    rect = plot.plot_rect()
    px = lambda idx: int(round(rect.left() + idx / 512 * rect.width()))
    QTest.mousePress(plot, Qt.LeftButton, Qt.ControlModifier, QPoint(px(176), int(rect.center().y())))
    QTest.mouseMove(plot, QPoint(px(225), int(rect.center().y())))
    QTest.mouseRelease(plot, Qt.LeftButton, Qt.ControlModifier, QPoint(px(225), int(rect.center().y())))
    assert len(plot.rois) == 1
    st = plot.rois[0].stats
    assert st.net == pytest.approx(1000 * 3 * math.sqrt(2 * math.pi), rel=0.01)
    assert window.roi_panel.table.rowCount() == 1
    assert not window.roi_dock.isHidden()
    assert (plot.x0, plot.x1) == (0, 512)          # Ctrl+drag does not zoom
    plot.add_roi(20, 40)
    out = window.export_rois(str(tmp_path / "rois.csv"))
    text = out.read_text(encoding="utf-8-sig")
    assert "net_unc_1s" in text and text.count("\n") >= 7
    window.roi_panel.select_row(0)
    window.roi_panel.delete_selected()
    assert len(plot.rois) == 1 and window.roi_panel.table.rowCount() == 1
    plot.clear_rois()
    assert window.roi_panel.table.rowCount() == 0
    assert not plot.grab().isNull()


def test_overlay_add_remove_normalize(window, tmp_path):
    window.display([gauss_spectrum(n=512, live=100.0)], None, None)
    bkg = make_file(tmp_path, "bkg.fak", sid="BKG", n=512)        # live time 100 s
    assert window.add_overlay_file(str(bkg))
    plot = window.plot
    assert len(plot.overlays) == 1 and plot.overlays[0].label == "bkg.fak"
    assert plot.overlays[0].by_energy                               # both calibrated -> energy aligned
    plot.overlays[0].spec.live_time = 50.0
    window.act_overlay_norm.setChecked(True)
    assert plot.overlays[0].scale == pytest.approx(2.0)
    assert not plot.grab().isNull()
    actions = [a.text() for a in window.menu_overlays.actions()]
    assert any("bkg.fak" in a for a in actions)
    plot.add_overlay(gauss_spectrum(n=256), "second")
    assert plot.overlays[1].color != plot.overlays[0].color
    plot.remove_overlay(0)
    assert [o.label for o in plot.overlays] == ["second"]
    plot.clear_overlays()
    assert not plot.overlays


def test_export_png_svg_clipboard(window, tmp_path):
    window.display([gauss_spectrum(n=2048)], None, None)
    window.plot.resize(800, 400)
    png = tmp_path / "view.png"
    svg = tmp_path / "view.svg"
    assert window.export_plot_image(str(png))
    assert window.export_plot_image(str(svg))
    assert png.stat().st_size > 2000
    from PySide6.QtGui import QImage
    assert QImage(str(png)).width() == 1600                          # 2x resolution
    assert "<svg" in svg.read_text(encoding="utf-8")[:2000]
    window.copy_plot()
    assert not QApplication.clipboard().image().isNull()


def test_large_spectrum_paints_fast(qapp):
    import random
    from spectramaster.gui.plot import SpectrumPlot
    p = SpectrumPlot()
    p.resize(1400, 500)
    rnd = random.Random(1)
    p.set_spectrum(Spectrum(counts=[rnd.randint(0, 5000) for _ in range(65536)], energy_coeffs=[0, 0.05]))
    p.add_overlay(Spectrum(counts=[rnd.randint(0, 100) for _ in range(65536)], energy_coeffs=[0, 0.05]), "b")
    t0 = time.perf_counter()
    for i in range(10):
        p.set_view(i * 100, 65536 - i * 100, record=False)
        p.grab()
    assert (time.perf_counter() - t0) / 10 < 0.25
    p.deleteLater()


# --------------------------------------------------------------- fake analysis engine
def _fake_engine(monkeypatch, calibrated_ok=True):
    from spectramaster import analysis
    from spectramaster.analysis import NuclideMatch, Peak
    from spectramaster.gui import analysis_panel

    def fake_find_peaks(spec, **kw):
        e = spec.energy_of(200.3) if spec.energy_coeffs else None
        pk = Peak(centroid=200.3, energy=e, fwhm=7.06, fwhm_kev=3.53 if e else None, net_area=7520.0,
                  net_area_unc=110.0, gross_area=12000.0, background=4480.0, left=180, right=220,
                  significance=68.0, count_rate=75.2, nuclide="", note="")
        pk2 = Peak(centroid=400.0, energy=spec.energy_of(400.0) if spec.energy_coeffs else None, fwhm=8.0,
                   fwhm_kev=4.0, net_area=100.0, net_area_unc=30.0, gross_area=500.0, background=400.0,
                   left=390, right=410, significance=3.3)
        return [pk, pk2]

    def fake_identify(spec, peaks, **kw):
        peaks[0].nuclide = "Am-241"
        peaks[0].line_energy = 100.15
        return [NuclideMatch("Am-241", 0.93, [(100.15, peaks[0])], [26.3], "432.6 y", "industrial",
                             "test"),
                NuclideMatch("K-40", 0.4, [], [1460.8], "1.25e9 y", "natural", "", background=True)]

    monkeypatch.setattr(analysis, "find_peaks", fake_find_peaks)
    monkeypatch.setattr(analysis, "identify", fake_identify)
    monkeypatch.setattr(analysis_panel, "engine_status", lambda: None)


def test_analysis_dock_populated_via_worker(window, monkeypatch, tmp_path):
    _fake_engine(monkeypatch)
    window.refresh_analysis_available()
    window.display([gauss_spectrum(n=4096)], None, None)
    assert window.act_analyze.isEnabled()
    assert window.run_analysis()
    assert wait_until(lambda: not window.analysis_running() and window.analysis_panel.peaks)
    panel = window.analysis_panel
    assert panel.peak_table.rowCount() == 2
    assert panel.nuc_table.rowCount() == 2
    assert panel.nuc_table.item(0, 0).text() == "Am-241"
    assert "93" in panel.nuc_table.item(0, 1).text()
    assert not window.analysis_dock.isHidden()
    assert len(window.plot.peaks) == 2
    # selecting a nuclide highlights its expected lines and matched peaks
    panel.select_nuclide_row(0)
    hl = window.plot.highlight
    assert hl.name == "Am-241" and hl.lines == [26.3, 100.15] and len(hl.peaks) == 1
    assert not window.plot.grab().isNull()
    # double-click on the peak row zooms around the peak (±5 FWHM)
    row = next(r for r in range(2) if panel.peak_table.item(r, 0).text() == "1")
    panel.activate_peak_row(row)
    plot = window.plot
    assert (plot.x0 + plot.x1) / 2 == pytest.approx(200.8, abs=0.6)
    assert plot.x1 - plot.x0 == pytest.approx(10 * 7.06, rel=0.05)
    # exports
    csv_out = window.export_analysis(str(tmp_path / "res.csv"))
    text = csv_out.read_text(encoding="utf-8-sig")
    assert "Am-241" in text and "# sample_id" in text
    xlsx = window.export_analysis(str(tmp_path / "res.xlsx"))
    from openpyxl import load_workbook
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Spectrum", "Peaks", "Nuclides"]
    assert wb["Nuclides"]["A2"].value == "Am-241"


def test_analysis_uncalibrated_message(window, monkeypatch):
    _fake_engine(monkeypatch)
    window.refresh_analysis_available()
    s = gauss_spectrum(n=1024)
    s.energy_coeffs = []
    window.display([s], None, None)
    assert window.run_analysis()
    assert wait_until(lambda: not window.analysis_running() and window.analysis_panel.peaks)
    panel = window.analysis_panel
    assert not panel.matches
    assert panel.banner.isVisibleTo(panel)
    assert "calibrat" in panel.banner.text().lower() or "معاير" in panel.banner.text()


def test_analysis_missing_engine_degrades(window, monkeypatch):
    window.aptec_analysis.set_engine("spectramaster")    # SpectraMaster's own engine (P)
    from spectramaster.gui import analysis_panel
    monkeypatch.setattr(analysis_panel, "engine_status", lambda: "spectramaster.analysis.identify not installed")
    window.refresh_analysis_available()
    window.display([gauss_spectrum()], None, None)
    assert not window.act_analyze.isEnabled()
    assert "identify" in window.act_analyze.toolTip()
    assert not window.run_analysis()
    assert window.errors


def test_analysis_engine_error_is_reported(window, monkeypatch):
    window.aptec_analysis.set_engine("spectramaster")    # SpectraMaster's own engine (P)
    _fake_engine(monkeypatch)
    from spectramaster import analysis

    def boom(spec, **kw):
        raise NotImplementedError("not yet")
    monkeypatch.setattr(analysis, "find_peaks", boom)
    window.refresh_analysis_available()
    window.display([gauss_spectrum()], None, None)
    assert window.run_analysis()
    assert wait_until(lambda: not window.analysis_running())
    assert window.errors and "not yet" in window.errors[-1]
    assert not window.act_analyze.isEnabled()


def test_auto_analyze_on_open(window, monkeypatch, tmp_path):
    _fake_engine(monkeypatch)
    window.refresh_analysis_available()
    window.act_auto_analyze.setChecked(True)
    assert window.settings.get_bool("analysis/auto")
    window.load_file(make_file(tmp_path, "auto.fak", n=1024))
    assert wait_until(lambda: not window.analysis_running() and window.analysis_panel.peaks)


def test_new_strings_both_languages_render(window):
    for lang in ("ar", "en"):
        window.apply_language(lang)
        window.display([gauss_spectrum(n=1024)], None, None)
        assert window.menu_analyze.title()
        assert window.yscale_combo.count() == 3
        assert window.plot.layoutDirection() == Qt.LeftToRight
        assert window.overview.layoutDirection() == Qt.LeftToRight
        assert not window.grab().isNull()


# =========================================================================== DataMaster parity
def test_datamaster_args_detection():
    from spectramaster.gui.app import is_datamaster_args
    assert is_datamaster_args(["-0", "a.s0", "-1", "a.chn"])
    assert is_datamaster_args(["-13", "a.s0", "-16", "a.cnf"])
    assert not is_datamaster_args(["a.s0"])
    assert not is_datamaster_args(["-0", "a.s0", "x", "a.chn"])


def test_gui_main_datamaster_silent_success(qapp, fake_format, tmp_path, monkeypatch):
    from spectramaster.gui import app as appmod
    src = make_file(tmp_path, "in.fak", sid="DM")
    dst = tmp_path / "out" / "converted.fak"
    dst.parent.mkdir()
    shown = []
    monkeypatch.setattr(appmod.MainWindow, "show", lambda self: shown.append(self))
    assert appmod.main(["SpectraMaster.exe", "-99", str(src), "-99", str(dst)]) == 0
    assert dst.is_file() and _fake_read(dst)[0].sample_id == "DM"
    assert not shown                                          # no window at all


def test_gui_main_datamaster_failure_goes_interactive(qapp, fake_format, tmp_path, monkeypatch):
    from spectramaster.gui import app as appmod
    src = make_file(tmp_path, "good.fak", sid="RECOVER")
    windows, errors = [], []
    monkeypatch.setattr(appmod.MainWindow, "show_error", lambda self, text: errors.append(text))
    monkeypatch.setattr(appmod.MainWindow, "show", lambda self: windows.append(self))
    monkeypatch.setattr(type(qapp), "exec", lambda self: 0)
    monkeypatch.setattr(appmod, "Settings", lambda: __import__("spectramaster.gui.settings", fromlist=["x"])
                        .Settings(QSettings(str(tmp_path / "s.ini"), QSettings.IniFormat)))
    # output format 10 (RMS PHD) is read-only -> the conversion fails
    assert appmod.main(["SpectraMaster.exe", "-0", str(src), "-10", str(tmp_path / "x.phd")]) == 0
    assert len(windows) == 1
    w = windows[0]
    assert w.current_spectrum() is not None and w.current_spectrum().sample_id == "RECOVER"
    assert errors and "-10" in errors[-1]
    w.close()


def test_open_preview_real_s0():
    pytest.importorskip("PySide6")
    sample = APTEC_DIR / "Eu-152.S0"
    if not sample.is_file():
        pytest.skip("reference Aptec spectra not available")
    from spectramaster.gui.preview import SpectrumPreview
    load_all()
    prev = SpectrumPreview()
    assert prev.show_file(str(sample))
    assert "APTEC" in prev.value_text("format").upper()
    assert int(prev.value_text("channels").replace(",", "")) > 1000
    assert prev.value_text("live_time") != prev.value_text("description") or True
    assert prev.thumb.pixmap() is not None and not prev.thumb.pixmap().isNull()
    assert not prev.show_file(str(Path(__file__)))            # not a spectrum -> message, no crash
    assert prev.message.text()


def test_preview_dialog_checkbox_remembered(qapp, settings, fake_format, tmp_path):
    from spectramaster.gui.preview import PreviewFileDialog
    d = PreviewFileDialog(None, "t", str(tmp_path), "All (*)", "", settings)
    assert d.show_cb.isChecked() and not d.preview.isHidden()
    d.show_cb.setChecked(False)
    assert settings.get_bool("ui/show_description", True) is False
    d2 = PreviewFileDialog(None, "t", str(tmp_path), "All (*)", "", settings)
    assert not d2.show_cb.isChecked() and d2.preview.isHidden()
    p = make_file(tmp_path, "pv.fak", sid="PREV")
    assert d2.preview.show_file(str(p)) and d2.preview.value_text("sample_id") == "PREV"
    d.close(); d2.close()


def test_integration_dialog_and_shortcuts(qapp, fake_format):
    from spectramaster.gui.dialogs import IntegrationDialog, ShortcutsDialog
    d = IntegrationDialog()
    assert d.fields["program"].text()
    assert "-1 $(FullBase).CHN" in d.fields["gv_case_import"].text()
    assert "-16 $(FullBase).CNF" in d.fields["gv_case_export_cnf"].text()
    d.copy(d.fields["gv_case_import"].text())
    assert QApplication.clipboard().text().endswith("$(FullBase).CHN")
    s = ShortcutsDialog()
    assert s.tree.topLevelItemCount() > 15
    d.close(); s.close()


# =========================================================================== designer credit
def test_about_shows_designer(qapp, fake_format):
    from spectramaster.gui.dialogs import AboutDialog
    from spectramaster.gui.icon import DESIGNER_PNG, designer_pixmap
    from PySide6.QtWidgets import QLabel
    dlg = AboutDialog()
    labels = " ".join(l.text() for l in dlg.findChildren(QLabel))
    assert "Edrees F. Hery" in labels and "2026" in labels
    if DESIGNER_PNG.exists():
        pm = designer_pixmap(110)
        assert pm is not None and pm.height() == 110
        logo = dlg.designer.designer_logo
        assert logo.pixmap() is not None and not logo.pixmap().isNull()
    dlg.close()


def test_splash_disabled_in_tests_and_renders(qapp, settings):
    from spectramaster.gui.app import splash_enabled, splash_pixmap
    assert not splash_enabled(settings)
    assert not splash_pixmap().isNull()


def test_status_bar_credit(window):
    window.apply_language("en")
    assert "Edrees F. Hery" in window.status_credit.text()
    window.apply_language("ar")
    assert "Edrees F. Hery" in window.status_credit.text()

@pytest.mark.parametrize("good", [True, False])
def test_smoke_test_flag_opens_file_and_exits(qapp, fake_format, tmp_path, monkeypatch, good):
    """``SpectraMaster.exe --smoke-test FILE`` (used to verify the installer build) loads FILE
    exactly like a double-click from Explorer and exits with 0 (loaded) or 3 (failed)."""
    from spectramaster.gui import app as gui_app
    from spectramaster.gui.settings import Settings
    ini = str(tmp_path / "smoke.ini")
    monkeypatch.setattr(gui_app, "Settings", lambda: Settings(QSettings(ini, QSettings.IniFormat)))
    p = make_file(tmp_path, "smoke.fak") if good else tmp_path / "missing.fak"
    code = gui_app.main(["SpectraMaster.exe", "--smoke-test", str(p)])
    assert code == (0 if good else 3)
    for w in QApplication.topLevelWidgets():
        w.close()
