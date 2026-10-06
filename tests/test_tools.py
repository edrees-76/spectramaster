"""Tests for spectramaster.tools (archive index/rescue, arithmetic, calibration)."""
from __future__ import annotations

import csv
import math
import random
import shutil
from dataclasses import replace
from datetime import datetime
from pathlib import Path

import numpy as np
import pytest

from spectramaster import cli
from spectramaster.convert import read_spectra
from spectramaster.formats import get_format, load_all
from spectramaster.model import Spectrum
from spectramaster.tools import (align_to, apply_calibration, apply_calibration_batch,
                                 calibrate_from_source, find_peaks, fit_energy_calibration,
                                 fit_fwhm_calibration, index_archive, parse_pairs, rebin,
                                 rescue_archive, same_calibration, scale_spectrum,
                                 subtract_background, sum_spectra, truncate, truncate_energy,
                                 verify_output)
from spectramaster.tools.arithmetic import SpectrumMathError

load_all()


def poisson_spec(n=1024, coeffs=(0.0, 0.5), seed=1, lt=100.0, peaks=(), bg=20.0, start=None, sid="S"):
    rng = np.random.default_rng(seed)
    ch = np.arange(n)
    mean = np.full(n, bg) * np.exp(-ch / (n * 1.5))
    for c, a, s in peaks:
        mean = mean + a / (s * math.sqrt(2 * math.pi)) * np.exp(-0.5 * ((ch - c) / s) ** 2)
    counts = rng.poisson(mean).astype(float)
    return Spectrum(counts=list(counts), energy_coeffs=list(coeffs), live_time=lt, real_time=lt * 1.02,
                    start_time=start or datetime(2009, 3, 6, 18, 51, 14), sample_id=sid,
                    description="عينة اختبار", detector="G4")


# ------------------------------------------------------------- arithmetic
def test_sum_same_calibration_exact():
    a, b = poisson_spec(seed=1), poisson_spec(seed=2, start=datetime(2008, 1, 1))
    s = sum_spectra([a, b], align="channel")
    assert s.counts == [x + y for x, y in zip(a.counts, b.counts)]
    assert s.live_time == 200.0 and s.real_time == pytest.approx(204.0)
    assert s.start_time == datetime(2008, 1, 1)
    assert a.counts == poisson_spec(seed=1).counts  # inputs untouched


def test_sum_energy_aligned_conserves_counts():
    a = poisson_spec(seed=1, coeffs=(0.0, 0.5))
    b = poisson_spec(seed=2, coeffs=(1.3, 0.47, 1e-6))     # different gain, inside a's range
    assert not same_calibration(a, b)
    with pytest.raises(SpectrumMathError):
        sum_spectra([a, b], align="channel")
    s = sum_spectra([a, b], align="auto")
    assert s.num_channels == a.num_channels and s.energy_coeffs == a.energy_coeffs
    outside = s.extra["counts_outside_range"]
    assert s.total_counts + outside == pytest.approx(a.total_counts + b.total_counts, rel=1e-12)
    assert outside < 0.01 * b.total_counts


def test_align_preserves_peak_energy():
    a = poisson_spec(coeffs=(0.0, 0.5), peaks=[(600, 50000, 2.0)], seed=3)   # 300 keV
    b = poisson_spec(coeffs=(0.0, 0.25), n=2048, peaks=[(1200, 50000, 4.0)], seed=4)
    al = align_to(b, a)
    i = int(np.argmax(al.counts))
    assert abs(a.energy_of(i) - 300.0) <= 0.5


def test_subtract_background_uncertainty():
    s = poisson_spec(seed=5, lt=100)
    b = poisson_spec(seed=6, lt=400)
    net, sigma = subtract_background(s, b)
    k = 0.25
    assert net.extra["background_scale"] == pytest.approx(k)
    for i in (0, 10, 500):
        assert net.counts[i] == pytest.approx(s.counts[i] - k * b.counts[i])
        assert sigma[i] == pytest.approx(math.sqrt(s.counts[i] + k * k * b.counts[i]))
    assert net.live_time == 100
    clipped, _ = subtract_background(s, b, scale=10, clip_negative=True)
    assert min(clipped.counts) == 0.0


def test_scale_rebin_truncate():
    s = poisson_spec(n=1001, coeffs=(2.0, 0.5, 1e-5), seed=7)
    s = replace(s, fwhm_coeffs=[1.0, 0.001])
    assert scale_spectrum(s, 2).total_counts == pytest.approx(2 * s.total_counts)
    for f in (2, 4, 3):
        r = rebin(s, f)
        assert r.total_counts == s.total_counts              # exact conservation
        assert r.num_channels == math.ceil(1001 / f)
        # new channel 10 covers old f*10 .. f*10+f-1, centre f*10+(f-1)/2
        assert r.energy_of(10) == pytest.approx(s.energy_of(f * 10 + (f - 1) / 2))
        assert r.fwhm_coeffs and np.polyval(r.fwhm_coeffs[::-1], 10) == pytest.approx(
            np.polyval(s.fwhm_coeffs[::-1], f * 10 + (f - 1) / 2))
    t = truncate(s, 100, 200)
    assert t.num_channels == 100 and t.first_channel == 100 and t.counts == s.counts[100:200]
    assert t.energies()[0] == pytest.approx(s.energy_of(100))
    t0 = truncate(s, 100, 200, renumber=True)
    assert t0.first_channel == 0 and t0.energy_of(0) == pytest.approx(s.energy_of(100))
    # rebin of a truncated spectrum with odd offset keeps energies right
    rt = rebin(truncate(s, 101, 301), 2)
    assert rt.energy_of(rt.first_channel) == pytest.approx(s.energy_of(101.5))
    te = truncate_energy(s, 100, 200)
    assert all(100 <= e <= 200 for e in te.energies())


# ------------------------------------------------------------- calibration
def test_energy_fit_recovery():
    true = [1.7, 0.3412, -7.3e-8]
    ch = [100, 900, 1933, 2700, 3434, 3902, 6000]
    pairs = [(c, true[0] + true[1] * c + true[2] * c * c) for c in ch]
    fit = fit_energy_calibration(pairs, order=2)
    assert fit.coeffs == pytest.approx(true, rel=1e-6, abs=1e-9)
    assert fit.max_abs_residual < 1e-8
    lin = fit_energy_calibration(pairs, order=1)
    assert lin.rms > 0 and "RMS residual" in lin.report()
    with pytest.raises(ValueError):
        fit_energy_calibration(pairs[:2], order=2)
    assert parse_pairs("1:2, 3:4;5=6") == [(1, 2), (3, 4), (5, 6)]
    fw = fit_fwhm_calibration([(100, 1.0), (1000, 1.45), (3000, 2.45)], order=1)
    assert fw.coeffs[1] == pytest.approx(0.0005, rel=1e-6)


def _eu_spectrum(coeffs=(1.2, 0.34, 0.0), n=8192, seed=11):
    from spectramaster.tools.calibrate import NUCLIDE_LINES
    c0, c1, c2 = coeffs
    pk = []
    for e, i in NUCLIDE_LINES["Eu-152"]:
        ch = (-c1 + math.sqrt(c1 * c1 - 4 * c2 * (c0 - e))) / (2 * c2) if c2 else (e - c0) / c1
        pk.append((ch, 3000 * i, 1.5 + ch / 4000))
    return poisson_spec(n=n, coeffs=[], peaks=pk, seed=seed, bg=40)


def test_find_peaks_and_calibrate_from_source():
    spec = _eu_spectrum()
    peaks = find_peaks(spec)
    assert len(peaks) >= 8
    sc = calibrate_from_source(spec, "Eu-152")
    assert len(sc.matches) >= 9
    for ch in (500, 2000, 4000):
        assert sc.fit.evaluate(ch) == pytest.approx(1.2 + 0.34 * ch, abs=0.15)
    assert sc.fit.max_abs_residual < 0.2
    assert sc.fwhm_fit is not None
    # with a (wrong but close) existing calibration the matcher is used
    spec2 = replace(spec, energy_coeffs=[3.0, 0.342])
    sc2 = calibrate_from_source(spec2, ["Eu-152"])
    assert sc2.method and sc2.fit.evaluate(4000) == pytest.approx(1361.2, abs=0.2)
    with pytest.raises(KeyError):
        calibrate_from_source(spec, "Xx-999")


def test_apply_calibration_batch(tmp_path):
    spec = _eu_spectrum()
    src = tmp_path / "in" / "eu.chn"
    src.parent.mkdir()
    get_format("ortec_chn").writer(spec, src)
    cal = poisson_spec()
    cal_path = tmp_path / "in" / "cal.chn"
    get_format("ortec_chn").writer(cal, cal_path)
    res = apply_calibration_batch([src, cal_path], tmp_path / "out", source="Eu-152", only_uncalibrated=True)
    assert res[0].error is None and res[0].output.exists()
    assert res[1].error.startswith("skipped")
    back = read_spectra(res[0].output)[0][0]
    assert back.energy_of(4000) == pytest.approx(1361.2, abs=0.3)
    new = apply_calibration(cal, [0, 1.0])
    assert new.extra["previous_energy_coeffs"] and cal.energy_coeffs == [0.0, 0.5]


# ------------------------------------------------------------- archive
def make_archive(root: Path) -> dict:
    files = {}
    s1 = poisson_spec(seed=21, sid="A1")
    s2 = poisson_spec(seed=22, sid="B2", start=datetime(2012, 5, 1, 8, 0, 0))
    s3 = replace(poisson_spec(seed=23, sid="C3"), energy_coeffs=[])
    (root / "01" / "tree").mkdir(parents=True)
    (root / "02" / "tree").mkdir(parents=True)
    files["a"] = root / "01" / "tree" / "a.S0"
    get_format("aptec_s0").writer(s1, files["a"])
    files["a_copy"] = root / "02" / "tree" / "a.S0"            # mirrored tree
    shutil.copy2(files["a"], files["a_copy"])
    files["b"] = root / "01" / "b.chn"
    get_format("ortec_chn").writer(s2, files["b"])
    files["b_spe"] = root / "02" / "b.spe"                    # same data, other format
    get_format("iaea_spe").writer(s2, files["b_spe"])
    files["c"] = root / "01" / "عينة.json"
    get_format("json").writer(s3, files["c"])
    files["bad"] = root / "01" / "corrupt.chn"
    files["bad"].write_bytes(bytes(random.Random(1).randrange(256) for _ in range(700)))
    (root / "01" / "report.docx").write_bytes(b"PK\x03\x04 not a spectrum")
    (root / "01" / "notes.dat").write_text("hello world\n")
    return files


def test_index_archive(tmp_path):
    root = tmp_path / "archive"
    files = make_archive(root)
    before = {p: p.read_bytes() for p in root.rglob("*") if p.is_file()}
    out = tmp_path / "cat" / "catalog.xlsx"
    seen = []
    res = index_archive(root, out, progress=lambda i, n, p: seen.append(i))
    assert seen and seen[-1] == res.files_scanned
    by_name = {e.file_name: e for e in res.entries}
    assert {"a.S0", "b.chn", "b.spe", "عينة.json", "corrupt.chn"} <= set(by_name)
    assert "report.docx" not in by_name and "notes.dat" not in by_name
    assert by_name["corrupt.chn"].error
    assert by_name["عينة.json"].status.startswith("warning") and "uncalibrated" in by_name["عينة.json"].status
    assert by_name["b.chn"].year == 2012 and by_name["b.chn"].sample_id == "B2"
    kinds = {g.kind: g for g in res.duplicates}
    assert len(kinds["identical file"].entries) == 2 and res.duplicate_files == 1
    assert {e.file_name for e in kinds["same spectrum data"].entries} == {"b.chn", "b.spe"}
    assert res.mirrored_folders[0][2] == 1
    # workbook layout
    from openpyxl import load_workbook
    wb = load_workbook(out)
    assert wb.sheetnames == ["Catalog", "Summary", "Duplicates", "Mirrored folders", "Errors"]
    ws = wb["Catalog"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws.cell(1, 1).value == "Path" and ws.max_row == len(res.entries) + 1
    assert any("عينة" in str(c.value) for c in ws["C"])
    assert wb["Errors"].max_row == 2
    # archive untouched, output inside archive refused
    assert before == {p: p.read_bytes() for p in root.rglob("*") if p.is_file()}
    with pytest.raises(ValueError):
        index_archive(root, root / "catalog.xlsx")


def _manifest(path):
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def test_rescue_verify_and_resume(tmp_path):
    root = tmp_path / "archive"
    make_archive(root)
    out = tmp_path / "rescued"
    res = rescue_archive(root, out, formats=["n42_2012", "ci_cnf"])
    st = {(e.source.name, e.format_key): e.status for e in res.entries}
    assert st[("a.S0", "n42_2012")].startswith("VERIFIED")
    assert st[("b.chn", "ci_cnf")].startswith("VERIFIED")
    assert st[("corrupt.chn", "")] == "ERROR"
    assert (out / "n42_2012" / "01" / "tree" / "a_S0.n42").exists()
    assert (out / "ci_cnf" / "02" / "tree" / "a_S0.cnf").exists()
    assert (out / "verification_report.xlsx").exists()
    rows = _manifest(out / "manifest.csv")
    assert len(rows) == len(res.entries) and all(r["source_sha256"] for r in rows)

    # resume: nothing redone
    res2 = rescue_archive(root, out, formats=["n42_2012", "ci_cnf"], resume=True)
    redone = [e for e in res2.entries if not e.skipped]
    assert all(e.source.name == "corrupt.chn" for e in redone)        # failures are retried
    assert sum(e.skipped for e in res2.entries) == len(res.entries) - 1
    # a damaged output is detected by its hash and redone
    target = out / "n42_2012" / "01" / "b_chn.n42"
    target.write_bytes(target.read_bytes() + b"\n")
    res3 = rescue_archive(root, out, formats=["n42_2012", "ci_cnf"], resume=True)
    redone = {(e.source.name, e.format_key) for e in res3.entries if not e.skipped}
    assert ("b.chn", "n42_2012") in redone and ("a.S0", "n42_2012") not in redone
    assert len(_manifest(out / "manifest.csv")) == len(res.entries)

    with pytest.raises(ValueError):
        rescue_archive(root, root / "out")


def test_rescue_detects_injected_corruption(tmp_path, monkeypatch):
    root = tmp_path / "archive"
    make_archive(root)
    fmt = get_format("n42_2012")
    real_writer = fmt.writer

    def corrupting(spec, path):
        c = list(spec.counts)
        c[7] += 1
        real_writer(replace(spec, counts=c, live_time=(spec.live_time or 0) + 5), path)
    monkeypatch.setattr(fmt, "writer", corrupting)
    res = rescue_archive(root, tmp_path / "out", formats=["n42_2012"])
    bad = [e for e in res.entries if e.format_key == "n42_2012"]
    assert bad and all(e.status == "FAILED" for e in bad)
    assert "counts differ" in bad[0].details and "live time" in bad[0].details


def test_verify_output_energy_refit_warning():
    s = poisson_spec(coeffs=(1.0, 0.34, 1e-7, 1e-11), n=8192)
    v = verify_output(s, replace(s))
    assert v.ok and v.status == "VERIFIED" and v.counts_equal
    o = replace(s, energy_coeffs=[1.0, 0.34, 1e-7])                  # cubic dropped: ~5.5 keV
    v = verify_output(s, o)
    assert not v.ok and "energy deviates" in v.problems[0]
    o = replace(s, energy_coeffs=[1.0, 0.34, 1e-7, 1.05e-11], description="x")
    v = verify_output(s, o)
    assert v.ok and v.warnings and v.status == "VERIFIED (warnings)"


# ------------------------------------------------------------- CLI
def test_cli_tools(tmp_path, capsys):
    root = tmp_path / "archive"
    files = make_archive(root)
    assert cli.main(["index", str(root), "-o", str(tmp_path / "c.xlsx"), "-q"]) == 0
    assert (tmp_path / "c.xlsx").exists()
    rc = cli.main(["rescue", str(root), "-o", str(tmp_path / "r"), "--to", "n42_2012", "-q"])
    assert rc == 1                                   # the corrupt file is reported
    assert "corrupt.chn" in capsys.readouterr().out
    assert cli.main(["rescue", str(root), "-o", str(tmp_path / "r"), "--to", "n42_2012", "--resume", "-q"]) == 1
    a, b = files["a"], files["a_copy"]
    assert cli.main(["sum", str(a), str(b), "-o", str(tmp_path / "sum.chn")]) == 0
    s = read_spectra(tmp_path / "sum.chn")[0][0]
    assert s.total_counts == 2 * read_spectra(a)[0][0].total_counts
    assert cli.main(["subtract", str(files["b"]), str(files["b_spe"]), "-o", str(tmp_path / "net.spe"),
                     "--sigma-csv", str(tmp_path / "sig.csv")]) == 0
    assert read_spectra(tmp_path / "net.spe")[0][0].total_counts == 0
    assert (tmp_path / "sig.csv").read_text().startswith("channel")
    assert cli.main(["rebin", str(files["b"]), "--factor", "4", "-o", str(tmp_path / "rb.chn")]) == 0
    assert read_spectra(tmp_path / "rb.chn")[0][0].num_channels == 256
    assert cli.main(["calibrate", str(files["c"]), "--pairs", "100:50,800:400", "-o", str(tmp_path / "cal"),
                     "--to", "spe"]) == 0
    out = capsys.readouterr().out
    assert "RMS residual" in out
    cal = read_spectra(next((tmp_path / "cal").glob("*.spe")))[0][0]
    assert cal.energy_coeffs[1] == pytest.approx(0.5)
    assert cli.main(["calibrate", str(files["c"])]) == 2
