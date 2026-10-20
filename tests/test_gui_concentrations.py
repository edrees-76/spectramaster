"""GUI: sample information and concentrations in the activity dialog (Release E, E5)."""
from __future__ import annotations

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from datetime import datetime, timedelta

import pytest

pytest.importorskip("PySide6")

from PySide6.QtCore import QSettings

from spectramaster.gui import i18n

START = datetime(2026, 5, 10, 8, 0)


@pytest.fixture(scope="module")
def qapp():
    from spectramaster.gui.app import create_app
    yield create_app(["test"])


@pytest.fixture()
def settings(tmp_path, monkeypatch):
    from spectramaster.gui.settings import Settings
    monkeypatch.setenv("SPECTRAMASTER_EFF_DIR", str(tmp_path / "curves"))
    return Settings(QSettings(str(tmp_path / "settings.ini"), QSettings.IniFormat))


@pytest.fixture(autouse=True)
def english():
    old = i18n.language()
    i18n.set_language("en")
    yield
    i18n.set_language(old)


@pytest.fixture(scope="module")
def spectrum():
    from spectramaster.validation.synthetic import HPGE, Continuum, TruePeak, make_spectrum
    s = make_spectrum(HPGE, [TruePeak(661.657, 2e4), TruePeak(364.49, 8e3)], Continuum(300, 400, 10), 7).spectrum
    s.start_time = START
    return s


@pytest.fixture()
def curve_path(tmp_path):
    from spectramaster.validation.benchmark import _flat_curve
    return _flat_curve().save(tmp_path / "curves" / "flat.json")


def _dialog(settings, spectrum, curve_path):
    from spectramaster.gui.quant_panel import ActivityDialog
    ad = ActivityDialog(None, spectrum, settings)
    ad.message = lambda text: setattr(ad, "last_error", text)
    ad.set_curve_path(curve_path)
    ad.nuc_edit.setText("Cs-137, I-131, Co-60")
    return ad


def _plain(text: str) -> str:
    return text.replace(i18n.LRI, "").replace(i18n.PDI, "")


def _row(ad, nuclide):
    t = ad.conc_view.table
    for r in range(t.rowCount()):
        if t.item(r, 0).text() == nuclide:
            return [_plain(t.item(r, c).text()) for c in range(t.columnCount())]
    raise AssertionError(nuclide)


def test_bq_per_kg_dry_label_and_values(qapp, settings, spectrum, curve_path):
    ad = _dialog(settings, spectrum, curve_path)
    sb = ad.sample_box
    sb.set_matrix("soil")
    assert sb.unit() == "Bq/kg"                         # proposed for soil
    assert not sb.mass_row[1].isHidden() and sb.volume_row[1].isHidden() and sb.filter_row[1].isHidden()
    sb.set_value("mass_kg", 0.5)
    sb.set_basis("dry")
    sb.set_value("quantity_rel_unc_pct", 1.0)
    res = ad.calculate()
    assert res is not None
    out = ad.concentration
    assert out["unit"] == "Bq/kg (dry)" and out["missing"] == []
    assert "Bq/kg (dry)" in _plain(ad.conc_view.title.text())
    assert "Bq/kg (dry)" in ad.conc_view.table.horizontalHeaderItem(1).text()
    cs = next(r for r in res.nuclides if r.nuclide == "Cs-137")
    row = next(r for r in out["rows"] if r["nuclide"] == "Cs-137")
    assert row["value"] == pytest.approx(cs.activity / 0.5)
    shown = _row(ad, "Cs-137")
    assert shown[1].startswith(f"{cs.activity / 0.5:.4g} ± ") and shown[3] == "detected"
    co = _row(ad, "Co-60")                              # not in the spectrum: < MDA, never zero
    assert co[1].startswith("< ") and co[3] == "< MDA"
    assert ad.sample_info.mass_kg == 0.5 and ad.sample_info.mass_basis == "dry"
    assert settings.get("quant/matrix") == "soil" and settings.get("quant/conc_unit") == "Bq/kg"
    # the old specific-activity column still works next to the concentration
    ad.qty_spin.setValue(0.5)
    ad.unit_combo.setCurrentText("kg")
    ad.calculate()
    assert "Bq/kg" in ad.table.horizontalHeaderItem(4).text()
    assert ad.concentration["unit"] == "Bq/kg (dry)"


def test_missing_mass_lists_the_missing_inputs_and_shows_no_value(qapp, settings, spectrum, curve_path):
    ad = _dialog(settings, spectrum, curve_path)
    ad.sample_box.set_matrix("food")
    ad.sample_box.set_basis("wet")
    assert ad.calculate() is not None
    out = ad.concentration
    assert out["rows"] == [] and out["missing"] == ["sample mass (kg)"]
    assert ad.conc_view.table.rowCount() == 0
    notes = ad.conc_view.notes_label.text()
    assert "Concentration not computed - missing: sample mass (kg)" in notes
    assert ad.table.rowCount() == 3                     # the activities are still shown
    ad.sample_box.set_basis(None)
    ad.calculate()
    assert "mass basis (dry or wet)" in ad.conc_view.notes_label.text()


def test_air_filter_with_sampling_times_corrects_i131_for_decay(qapp, settings, spectrum, curve_path):
    from spectramaster.analysis.activity import sampling_decay_factor
    ad = _dialog(settings, spectrum, curve_path)
    sb = ad.sample_box
    sb.set_matrix("air_filter")
    assert sb.unit() == "Bq/m3" and not sb.filter_row[1].isHidden() and sb.mass_row[1].isHidden()
    sb.set_sampling(START - timedelta(days=7), START)
    sb.set_value("flow_rate_m3_h", 10.0)
    sb.set_value("filter_fraction", 1.0)
    sb.set_value("collection_efficiency", 0.95)
    sb.eff_note.setText("manufacturer value")
    res = ad.calculate()
    assert res is not None and ad.concentration["missing"] == []
    i131 = next(r for r in ad.concentration["rows"] if r["nuclide"] == "I-131")
    cs = next(r for r in ad.concentration["rows"] if r["nuclide"] == "Cs-137")
    assert i131["factor"] > 1.3
    assert i131["factor"] == pytest.approx(sampling_decay_factor(8.0252 * 86400, 7 * 86400), rel=1e-3)
    assert cs["factor"] == pytest.approx(1.0, abs=1e-3)
    a = next(r for r in res.nuclides if r.nuclide == "Cs-137")
    assert cs["value"] == pytest.approx(a.activity / (1680.0 * 0.95), rel=1e-3)
    assert float(_row(ad, "I-131")[4]) > 1.3
    notes = ad.conc_view.notes_label.text()
    assert "assumption: manufacturer value" in notes and "Decay during sampling corrected" in notes
    assert "Bq/m³" in _plain(ad.conc_view.title.text())


def test_air_filter_missing_inputs(qapp, settings, spectrum, curve_path):
    ad = _dialog(settings, spectrum, curve_path)
    ad.sample_box.set_matrix("air_filter")
    ad.calculate()
    assert ad.concentration["rows"] == [] and len(ad.concentration["missing"]) == 3
    notes = ad.conc_view.notes_label.text()
    assert "measured filter fraction" in notes and "collection efficiency" in notes


def test_invalid_number_does_not_block_the_activities(qapp, settings, spectrum, curve_path):
    ad = _dialog(settings, spectrum, curve_path)
    ad.sample_box.set_unit("Bq/L")
    ad.sample_box._num["volume_l"].setText("abc")
    assert ad.calculate() is not None and ad.table.rowCount() == 3
    assert ad.concentration is None and "'abc' is not a number" in ad.conc_view.notes_label.text()
    ad.sample_box._num["volume_l"].setText("2")
    ad.calculate()
    assert ad.concentration["unit"] == "Bq/L" and ad.concentration["rows"]


def test_export_contains_the_concentration(qapp, settings, spectrum, curve_path, tmp_path):
    from openpyxl import load_workbook
    ad = _dialog(settings, spectrum, curve_path)
    ad.sample_box.set_matrix("soil")
    ad.sample_box.set_value("mass_kg", 0.5)
    ad.sample_box.set_basis("dry")
    ad.calculate()
    p = ad.export(str(tmp_path / "out.xlsx"))
    ws = load_workbook(p)["Concentration"]
    assert ws["B1"].value == "Bq/kg (dry)"
    csv = ad.export(str(tmp_path / "out.csv"))
    assert "Bq/kg (dry)" in csv.read_text(encoding="utf-8-sig")


def test_arabic_dialog_renders_concentrations(qapp, settings, spectrum, curve_path):
    i18n.set_language("ar")
    ad = _dialog(settings, spectrum, curve_path)
    ar = i18n.STRINGS["ar"]
    assert ad.sample_box.title() == ar["quant_conc_group"]
    ad.sample_box.set_matrix("soil")
    ad.sample_box.set_value("mass_kg", 0.5)
    ad.sample_box.set_basis("dry")
    assert ad.calculate() is not None
    title = _plain(ad.conc_view.title.text())
    assert title.startswith("التركيز") and "Bq/kg (جاف)" in title
    assert _row(ad, "Co-60")[3] == ar["quant_res_lt_mda"]
    ad.sample_box.set_value("mass_kg", None)
    ad.calculate()
    assert ar["quant_conc_miss_mass"] in ad.conc_view.notes_label.text()
    assert not ad.grab().isNull()


def test_every_concentration_string_has_arabic():
    from spectramaster.gui import quant_strings as qs
    keys = [k for k in qs.EN if k.startswith("conc_")]
    assert keys and all(k in qs.AR for k in keys)


def test_calculate_on_a_worker_thread(qapp, settings, spectrum, curve_path):
    import time
    from PySide6.QtWidgets import QApplication
    ad = _dialog(settings, spectrum, curve_path)
    expected = ad.calculate()
    ad.results = None
    assert ad.calculate(threaded=True) is None
    assert ad.busy and not ad.btn_calc.isEnabled()
    assert ad.calculate(threaded=True) is None             # a second click while busy does nothing
    end = time.monotonic() + 60
    while ad.busy and time.monotonic() < end:
        QApplication.processEvents()
        time.sleep(0.01)
    QApplication.processEvents()
    assert not ad.busy and ad.btn_calc.isEnabled()
    got = {r.nuclide: r.activity for r in ad.results.nuclides}
    assert got == {r.nuclide: r.activity for r in expected.nuclides}
    assert ad.table.rowCount() == len(expected.nuclides)
