"""Detector quality control (QC): control charts for check-source / background measurements.

ISO/IEC 17025 (7.7 "ensuring the validity of results") asks a laboratory to monitor
its measuring systems with control charts.  For a gamma spectrometer the usual
records are, per detector and per check source:

* **centroid energy** of the monitored lines   -> gain / offset drift
* **FWHM** (keV)                                 -> resolution
* **net count rate**, decay-corrected to a reference date -> efficiency stability
* **background count rate** in an energy window  -> contamination / shielding
* live / real / dead time

Model
-----
A :class:`QCTrack` is one detector + one check source (or a background
measurement) with its monitored lines, tolerances and the list of
:class:`QCMeasurement` records.  A measurement stores the *raw* peak list of the
spectrum (``analysis.find_peaks``) so lines can be added / tolerances changed later
without re-reading the spectrum files; :func:`evaluate` derives the series,
statistics and rule violations on demand.

Statistics
----------
Baseline = the first ``baseline_n`` non-excluded points, or the points in
``[baseline_start, baseline_end]``.  Mean and sample SD give Shewhart limits
(warning +/-2 sigma, action +/-3 sigma).  Westgard rules (evaluated on all
non-excluded points in date order, z = (x - mean) / SD):

=======  ====================================================================
1-2s     |z| > 2 (warning only)
1-3s     |z| > 3
2-2s     two consecutive points beyond 2 sigma on the same side
R-4s     two consecutive points beyond 2 sigma on opposite sides (range > 4 sigma)
4-1s     four consecutive points beyond 1 sigma on the same side
10-x     ten consecutive points on the same side of the mean
SPLIT    line found split into 2+ peaks (gain/zero jump during the count)
TOL      fixed tolerance violated (centroid +/- keV of nominal, FWHM max, rate / background +/- % of baseline mean)
=======  ====================================================================

Storage: one JSON file per track in :func:`default_qc_dir`
(``%APPDATA%\\SpectraMaster\\qc``), written atomically.

Public API
----------
``QCTrack``, ``MonitoredLine``, ``QCMeasurement``, ``QCStore``
``preset_lines(source)``                       monitored lines of a check source
``analyse_spectrum(spectrum, track, path="", index=0) -> QCMeasurement``
``add_spectrum(track, spectrum, path="", index=0) -> QCMeasurement | None`` (None = duplicate)
``import_paths(track, paths, recursive=True, progress=None) -> ImportResult``
``reanalyse(track, progress=None) -> ImportResult``  re-read every source file
``evaluate(track) -> QCEvaluation``           series + statistics + violations
``export_track_xlsx(track, path, ev=None)``
"""
from __future__ import annotations

import json
import math
import os
import uuid
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Optional, Sequence

from ..model import Spectrum

__all__ = [
    "MonitoredLine", "QCMeasurement", "QCTrack", "QCStore", "Series", "SeriesPoint", "Violation",
    "QCEvaluation", "ImportResult", "ViolationRun", "violation_runs", "SOURCE_PRESETS", "WESTGARD_RULES", "DEFAULT_RULES", "METRICS",
    "default_qc_dir", "preset_lines", "half_life_of", "decay_factor", "analyse_spectrum",
    "add_spectrum", "import_paths", "reanalyse", "evaluate", "export_track_xlsx",
]

# ------------------------------------------------------------------ constants
WESTGARD_RULES = ("1-3s", "2-2s", "R-4s", "4-1s", "10-x")
DEFAULT_RULES = list(WESTGARD_RULES)
WARNING_RULES = {"1-2s"}
METRICS = ("centroid", "fwhm", "rate", "background")
METRIC_UNITS = {"centroid": "keV", "fwhm": "keV", "rate": "cps", "background": "cps"}

STATUS_OK, STATUS_WARN, STATUS_FAIL, STATUS_EXCLUDED, STATUS_NA = "OK", "WARN", "FAIL", "EXCL", "N/A"

# check source -> [(nuclide, keV, default FWHM max keV or None)]
SOURCE_PRESETS: dict[str, list[tuple[str, float, Optional[float]]]] = {
    "Cs-137": [("Cs-137", 661.657, None)],
    "Co-60": [("Co-60", 1173.228, None), ("Co-60", 1332.492, 2.0)],
    "Cs-137 + Co-60": [("Cs-137", 661.657, None), ("Co-60", 1173.228, None), ("Co-60", 1332.492, 2.0)],
    "Eu-152": [("Eu-152", 121.782, None), ("Eu-152", 344.279, None), ("Eu-152", 778.904, None),
               ("Eu-152", 1408.013, None)],
    "Ba-133": [("Ba-133", 80.998, None), ("Ba-133", 356.013, None)],
    "Am-241": [("Am-241", 59.541, None)],
    "Na-22": [("Na-22", 1274.537, None)],
    "Background": [("K-40", 1460.822, None), ("Th-232", 2614.511, None)],
}

_SECONDS_PER_DAY = 86400.0


def default_qc_dir() -> Path:
    """``%APPDATA%\\SpectraMaster\\qc`` (``~/.spectramaster/qc`` elsewhere)."""
    base = os.environ.get("APPDATA")
    if base:
        return Path(base) / "SpectraMaster" / "qc"
    return Path.home() / ".spectramaster" / "qc"


def half_life_of(nuclide: str) -> Optional[float]:
    """Half-life in seconds from the bundled nuclide library (None if unknown)."""
    if not nuclide:
        return None
    try:
        from ..analysis.library import load_library
        lib = load_library()
    except Exception:
        return None
    key = nuclide.strip()
    nd = lib.get(key)
    if nd is None:
        norm = key.lower().replace("-", "").replace(" ", "")
        for name, d in lib.items():
            if name.lower().replace("-", "").replace(" ", "") == norm:
                nd = d
                break
    return nd.half_life_s if nd is not None else None


def intensity_of(nuclide: str, energy: float, tol: float = 0.6) -> Optional[float]:
    """Emission probability (%) of the library line closest to `energy` (within `tol` keV)."""
    try:
        from ..analysis.library import load_library
        nd = load_library().get(nuclide)
    except Exception:
        return None
    if nd is None:
        return None
    best = min((l for l in nd.lines if l.kind == "g"), key=lambda l: abs(l.energy - energy), default=None)
    return best.intensity if best is not None and abs(best.energy - energy) <= tol else None


def decay_factor(half_life_s: Optional[float], t: Optional[datetime], ref: Optional[datetime],
                 real_time: Optional[float] = None) -> float:
    """Factor that converts a count rate measured from `t` (lasting `real_time` s)
    to the rate at the reference date `ref`: ``exp(lambda*(t-ref)) * lambda*Tr/(1-exp(-lambda*Tr))``."""
    if not half_life_s or half_life_s <= 0 or not math.isfinite(half_life_s) or t is None or ref is None:
        return 1.0
    lam = math.log(2.0) / half_life_s
    f = math.exp(lam * (t - ref).total_seconds())
    if real_time and real_time > 0:
        x = lam * real_time
        if x > 1e-9:
            f *= x / (1.0 - math.exp(-x))
    return f


def _poly(coeffs: Sequence[float], x: float) -> float:
    out = 0.0
    for c in reversed(list(coeffs)):
        out = out * x + c
    return out


def _dpoly(coeffs: Sequence[float], x: float) -> float:
    return _poly([k * c for k, c in enumerate(coeffs)][1:], x) if len(coeffs) > 1 else 0.0


def _dt(s) -> Optional[datetime]:
    if s is None or s == "":
        return None
    if isinstance(s, datetime):
        return s
    return datetime.fromisoformat(str(s))


def _iso(d: Optional[datetime]) -> Optional[str]:
    return d.isoformat(timespec="seconds") if d else None


# ------------------------------------------------------------------ model
@dataclass
class MonitoredLine:
    energy: float                       # nominal keV
    nuclide: str = ""
    fwhm_max: Optional[float] = None    # tolerance: FWHM must be <= this (keV)
    half_life_s: Optional[float] = None  # override; default from the nuclide library
    intensity: Optional[float] = None   # emission probability %, default from the library

    @property
    def key(self) -> str:
        return f"{self.energy:.3f}"

    @property
    def label(self) -> str:
        return f"{self.nuclide} {self.energy:.1f} keV".strip()

    def half_life(self) -> Optional[float]:
        return self.half_life_s if self.half_life_s else half_life_of(self.nuclide)

    def emission(self) -> Optional[float]:
        return self.intensity if self.intensity else intensity_of(self.nuclide, self.energy)


@dataclass
class QCMeasurement:
    id: str                             # spectrum fingerprint (content hash)
    date: Optional[datetime]
    source_path: str = ""
    spectrum_index: int = 0
    sample_id: str = ""
    live_time: Optional[float] = None
    real_time: Optional[float] = None
    energy_coeffs: list[float] = field(default_factory=list)
    total_counts: float = 0.0
    window_counts: Optional[float] = None
    window: tuple[Optional[float], Optional[float]] = (None, None)
    peaks: list[dict] = field(default_factory=list)   # {ch, e, fw_ch, fw_kev, net, unc, sig}
    excluded: bool = False
    note: str = ""
    added: Optional[datetime] = None

    @property
    def dead_time_pct(self) -> Optional[float]:
        if self.live_time and self.real_time and self.real_time > 0:
            return 100.0 * (self.real_time - self.live_time) / self.real_time
        return None

    @property
    def window_rate(self) -> Optional[float]:
        if self.window_counts is None or not self.live_time:
            return None
        return self.window_counts / self.live_time

    @property
    def file_name(self) -> str:
        return Path(self.source_path).name if self.source_path else ""

    def to_json(self) -> dict:
        d = asdict(self)
        d["date"] = _iso(self.date)
        d["added"] = _iso(self.added)
        d["window"] = list(self.window)
        return d

    @classmethod
    def from_json(cls, d: dict) -> "QCMeasurement":
        d = dict(d)
        d["date"] = _dt(d.get("date"))
        d["added"] = _dt(d.get("added"))
        d["window"] = tuple(d.get("window") or (None, None))
        known = {f for f in cls.__dataclass_fields__}
        return cls(**{k: v for k, v in d.items() if k in known})


@dataclass
class QCTrack:
    name: str
    id: str = field(default_factory=lambda: uuid.uuid4().hex[:12])
    detector: str = ""
    detector_type: str = "HPGe"         # HPGe / NaI / LaBr3 / CZT
    source: str = ""                    # check source, e.g. "Cs-137 + Co-60", or "Background"
    kind: str = "source"                # "source" or "background" (no decay correction)
    lines: list[MonitoredLine] = field(default_factory=list)
    reference_date: Optional[datetime] = None          # decay-correction date (None = first measurement)
    reference_activity: dict[str, float] = field(default_factory=dict)   # nuclide -> Bq at reference_date
    match_tol_kev: float = 2.0          # peak <-> line matching tolerance
    centroid_mode: str = "file"         # "file": energy from each file's calibration; "fixed": reference_coeffs
    reference_coeffs: list[float] = field(default_factory=list)          # fixed calibration (auto = first measurement)
    bkg_window: tuple[Optional[float], Optional[float]] = (None, None)    # keV; None = spectrum end
    baseline_n: int = 20
    baseline_start: Optional[datetime] = None
    baseline_end: Optional[datetime] = None
    rules: list[str] = field(default_factory=lambda: list(DEFAULT_RULES))
    centroid_tol_kev: Optional[float] = 0.5
    rate_tol_pct: Optional[float] = 5.0
    bkg_tol_pct: Optional[float] = None
    notes: str = ""
    created: Optional[datetime] = field(default_factory=lambda: datetime.now().replace(microsecond=0))
    measurements: list[QCMeasurement] = field(default_factory=list)

    # ------------------------------------------------------------ helpers
    def sorted_measurements(self) -> list[QCMeasurement]:
        return sorted(self.measurements, key=lambda m: (m.date or datetime.max, m.source_path, m.spectrum_index))

    def measurement(self, mid: str) -> Optional[QCMeasurement]:
        for m in self.measurements:
            if m.id == mid:
                return m
        return None

    def line(self, key_or_energy) -> Optional[MonitoredLine]:
        for l in self.lines:
            if l.key == key_or_energy or (isinstance(key_or_energy, (int, float))
                                          and abs(l.energy - float(key_or_energy)) < 0.05):
                return l
        return None

    def effective_reference_date(self) -> Optional[datetime]:
        if self.reference_date:
            return self.reference_date
        ds = [m.date for m in self.measurements if m.date and not m.excluded] or \
             [m.date for m in self.measurements if m.date]
        return min(ds) if ds else None

    def effective_coeffs(self, m: QCMeasurement) -> list[float]:
        """Calibration used to convert channel -> keV for measurement `m`."""
        if self.centroid_mode == "fixed":
            if self.reference_coeffs:
                return list(self.reference_coeffs)
            first = next((x for x in self.sorted_measurements() if x.energy_coeffs), None)
            if first is not None:
                return list(first.energy_coeffs)
        return list(m.energy_coeffs or self.reference_coeffs)

    def set_lines_from_preset(self, source: str) -> None:
        self.source = source
        self.lines = preset_lines(source)
        if source.lower().startswith("background"):
            self.kind = "background"

    # ------------------------------------------------------------ json
    def to_json(self) -> dict:
        return {
            "format": "spectramaster-qc-track", "version": 1,
            "id": self.id, "name": self.name, "detector": self.detector,
            "detector_type": self.detector_type, "source": self.source, "kind": self.kind,
            "lines": [asdict(l) for l in self.lines],
            "reference_date": _iso(self.reference_date),
            "reference_activity": dict(self.reference_activity),
            "match_tol_kev": self.match_tol_kev, "centroid_mode": self.centroid_mode,
            "reference_coeffs": list(self.reference_coeffs), "bkg_window": list(self.bkg_window),
            "baseline_n": self.baseline_n, "baseline_start": _iso(self.baseline_start),
            "baseline_end": _iso(self.baseline_end), "rules": list(self.rules),
            "centroid_tol_kev": self.centroid_tol_kev, "rate_tol_pct": self.rate_tol_pct,
            "bkg_tol_pct": self.bkg_tol_pct, "notes": self.notes, "created": _iso(self.created),
            "measurements": [m.to_json() for m in self.sorted_measurements()],
        }

    @classmethod
    def from_json(cls, d: dict) -> "QCTrack":
        t = cls(name=d.get("name", ""), id=d.get("id") or uuid.uuid4().hex[:12])
        t.detector = d.get("detector", "")
        t.detector_type = d.get("detector_type", "HPGe")
        t.source = d.get("source", "")
        t.kind = d.get("kind", "source")
        lf = set(MonitoredLine.__dataclass_fields__)
        t.lines = [MonitoredLine(**{k: v for k, v in l.items() if k in lf}) for l in d.get("lines", [])]
        t.reference_date = _dt(d.get("reference_date"))
        t.reference_activity = {str(k): float(v) for k, v in (d.get("reference_activity") or {}).items()}
        t.match_tol_kev = float(d.get("match_tol_kev", 2.0))
        t.centroid_mode = d.get("centroid_mode", "file")
        t.reference_coeffs = [float(c) for c in d.get("reference_coeffs", [])]
        t.bkg_window = tuple(d.get("bkg_window") or (None, None))
        t.baseline_n = int(d.get("baseline_n", 20))
        t.baseline_start = _dt(d.get("baseline_start"))
        t.baseline_end = _dt(d.get("baseline_end"))
        t.rules = list(d.get("rules", DEFAULT_RULES))
        t.centroid_tol_kev = d.get("centroid_tol_kev")
        t.rate_tol_pct = d.get("rate_tol_pct")
        t.bkg_tol_pct = d.get("bkg_tol_pct")
        t.notes = d.get("notes", "")
        t.created = _dt(d.get("created"))
        t.measurements = [QCMeasurement.from_json(m) for m in d.get("measurements", [])]
        return t


def preset_lines(source: str) -> list[MonitoredLine]:
    """Monitored lines of a check source preset ("Cs-137", "Co-60", "Cs-137 + Co-60",
    "Eu-152", "Background" ...) or of any comma/plus separated nuclide list."""
    if source in SOURCE_PRESETS:
        return [MonitoredLine(e, n, fm) for n, e, fm in SOURCE_PRESETS[source]]
    out: list[MonitoredLine] = []
    from .calibrate import NUCLIDE_LINES
    for raw in source.replace("+", ",").replace(";", ",").split(","):
        name = raw.strip()
        if not name:
            continue
        if name in SOURCE_PRESETS:
            out.extend(MonitoredLine(e, n, fm) for n, e, fm in SOURCE_PRESETS[name])
            continue
        lines = NUCLIDE_LINES.get(name)
        if lines:
            strongest = sorted(lines, key=lambda t: -t[1])[:3]
            out.extend(MonitoredLine(e, name) for e, _i in sorted(strongest))
    return sorted(out, key=lambda l: l.energy)


# ------------------------------------------------------------------ storage
class QCStore:
    """One ``<id>.json`` per track under `root` (default :func:`default_qc_dir`)."""

    def __init__(self, root=None):
        self.root = Path(root) if root else default_qc_dir()

    def path_of(self, track: QCTrack) -> Path:
        return self.root / f"{track.id}.json"

    def list_tracks(self) -> list[QCTrack]:
        out = []
        if self.root.is_dir():
            for p in sorted(self.root.glob("*.json")):
                try:
                    out.append(self.load(p))
                except Exception:
                    continue
        return sorted(out, key=lambda t: t.name.lower())

    def load(self, path_or_id) -> QCTrack:
        p = Path(path_or_id)
        if not p.suffix:
            p = self.root / f"{path_or_id}.json"
        return QCTrack.from_json(json.loads(p.read_text(encoding="utf-8")))

    def save(self, track: QCTrack) -> Path:
        self.root.mkdir(parents=True, exist_ok=True)
        p = self.path_of(track)
        tmp = p.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(track.to_json(), ensure_ascii=False, indent=1), encoding="utf-8")
        os.replace(tmp, p)
        return p

    def delete(self, track: QCTrack) -> None:
        p = self.path_of(track)
        if p.exists():
            p.unlink()


# ------------------------------------------------------------------ analysis
def _fingerprint(s: Spectrum) -> str:
    from .archive import spectrum_fingerprint
    return spectrum_fingerprint(s)[:24]


def _window_counts(s: Spectrum, window) -> Optional[float]:
    lo, hi = window if window else (None, None)
    if lo is None and hi is None:
        return float(sum(s.counts))
    if not s.energy_coeffs:
        return None
    tot = 0.0
    for i, c in enumerate(s.counts):
        e = _poly(s.energy_coeffs, s.first_channel + i)
        if (lo is None or e >= lo) and (hi is None or e <= hi):
            tot += c
    return tot


def analyse_spectrum(spectrum: Spectrum, track: Optional[QCTrack] = None, path: str = "",
                     index: int = 0) -> QCMeasurement:
    """Peak search on one QC spectrum -> a measurement record (not yet added to the track)."""
    from .. import analysis
    fp = analysis.find_peaks
    if not callable(fp):          # package attribute shadowed by the submodule
        fp = fp.find_peaks
    peaks = []
    for p in fp(spectrum) or []:
        e = p.energy
        if e is None and track is not None and track.reference_coeffs:
            e = _poly(track.reference_coeffs, p.centroid)
        peaks.append({"ch": round(float(p.centroid), 4), "e": None if e is None else round(float(e), 4),
                      "fw_ch": round(float(p.fwhm), 4),
                      "fw_kev": None if p.fwhm_kev is None else round(float(p.fwhm_kev), 4),
                      "net": round(float(p.net_area), 2), "unc": round(float(p.net_area_unc), 2),
                      "sig": round(float(p.significance), 2)})
    window = track.bkg_window if track is not None else (None, None)
    return QCMeasurement(
        id=_fingerprint(spectrum), date=spectrum.start_time,
        source_path=path or spectrum.source_path or "", spectrum_index=index,
        sample_id=spectrum.sample_id or "", live_time=spectrum.live_time, real_time=spectrum.real_time,
        energy_coeffs=[float(c) for c in spectrum.energy_coeffs], total_counts=float(sum(spectrum.counts)),
        window_counts=_window_counts(spectrum, window), window=tuple(window), peaks=peaks,
        added=datetime.now().replace(microsecond=0))


def add_spectrum(track: QCTrack, spectrum: Spectrum, path: str = "", index: int = 0
                 ) -> Optional[QCMeasurement]:
    """Analyse and append; returns None when the same spectrum is already in the track."""
    fp = _fingerprint(spectrum)
    if any(m.id == fp for m in track.measurements):
        return None
    m = analyse_spectrum(spectrum, track, path, index)
    track.measurements.append(m)
    return m


@dataclass
class ImportResult:
    added: list[QCMeasurement] = field(default_factory=list)
    duplicates: list[str] = field(default_factory=list)
    errors: list[tuple[str, str]] = field(default_factory=list)
    cancelled: bool = False


ProgressCB = Callable[[int, int, str], Optional[bool]]


def _expand(paths: Iterable, recursive: bool) -> list[Path]:
    from .archive import find_spectrum_files
    out: list[Path] = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            out.extend(find_spectrum_files(p, recursive=recursive))
        elif p.exists():
            out.append(p)
    return out


def import_paths(track: QCTrack, paths: Iterable, recursive: bool = True,
                 progress: Optional[ProgressCB] = None) -> ImportResult:
    """Import QC spectra (files and/or folders).  Files are only read, never written.
    Records end up sorted by acquisition date."""
    from ..convert import read_spectra
    files = _expand(paths, recursive)
    res = ImportResult()
    for i, f in enumerate(files):
        if progress is not None and progress(i, len(files), str(f)) is False:
            res.cancelled = True
            break
        try:
            spectra, _fmt = read_spectra(f)
        except Exception as exc:
            res.errors.append((str(f), f"{type(exc).__name__}: {exc}"))
            continue
        for k, s in enumerate(spectra):
            try:
                m = add_spectrum(track, s, str(f), k)
            except Exception as exc:
                res.errors.append((str(f), f"{type(exc).__name__}: {exc}"))
                continue
            if m is None:
                res.duplicates.append(str(f))
            else:
                res.added.append(m)
    if progress is not None and not res.cancelled:
        progress(len(files), len(files), "")
    track.measurements = track.sorted_measurements()
    return res


def reanalyse(track: QCTrack, progress: Optional[ProgressCB] = None) -> ImportResult:
    """Re-read every measurement's source file and redo the peak search (keeps the
    excluded flag and note).  Needed after changing the background window."""
    from ..convert import read_spectra
    res = ImportResult()
    ms = track.sorted_measurements()
    for i, m in enumerate(ms):
        if progress is not None and progress(i, len(ms), m.source_path) is False:
            res.cancelled = True
            break
        try:
            spectra, _ = read_spectra(Path(m.source_path))
            s = spectra[m.spectrum_index]
            new = analyse_spectrum(s, track, m.source_path, m.spectrum_index)
        except Exception as exc:
            res.errors.append((m.source_path, f"{type(exc).__name__}: {exc}"))
            continue
        new.excluded, new.note, new.added = m.excluded, m.note, m.added
        track.measurements[track.measurements.index(m)] = new
        res.added.append(new)
    return res


# ------------------------------------------------------------------ evaluation
@dataclass
class LineResult:
    centroid_ch: float
    centroid_kev: Optional[float]
    fwhm_kev: Optional[float]
    net: float
    net_unc: float
    rate: Optional[float]           # cps (live time)
    rate_unc: Optional[float]
    rate_corr: Optional[float]      # decay-corrected to the reference date
    decay: float = 1.0
    efficiency: Optional[float] = None   # rate_corr / (A_ref * I)
    components: int = 1             # >1: the line was found split into several peaks


def line_result(track: QCTrack, m: QCMeasurement, line: MonitoredLine,
                ref_date: Optional[datetime] = None, coeffs: Optional[list[float]] = None
                ) -> Optional[LineResult]:
    """Match `line` to the stored peaks within ``match_tol_kev``.

    A line that the peak fitter split into several components (peak doubling after a
    zero / gain jump during the acquisition, a badly degraded detector) is recombined:
    net areas are summed, the centroid is area-weighted and the FWHM is that of the
    combined distribution (sigma^2 = sum w*(sigma_i^2 + (c_i - c)^2)), so the QC
    values reflect the real line shape; ``LineResult.components`` > 1 marks it."""
    coeffs = coeffs if coeffs is not None else track.effective_coeffs(m)
    found = []
    for p in m.peaks:
        e = _poly(coeffs, p["ch"]) if coeffs else p.get("e")
        if e is None:
            continue
        # match with the file's own calibration when it has one (a fixed reference
        # calibration may be several keV off after a gain drift), report with `coeffs`
        e_match = p.get("e") if p.get("e") is not None else e
        if abs(e_match - line.energy) <= track.match_tol_kev:
            found.append(p)
    if not found:
        return None
    pos = [q for q in found if q["net"] > 0] or found
    wsum = sum(max(q["net"], 0.0) for q in pos) or float(len(pos))
    w = [(max(q["net"], 0.0) / wsum) if wsum else 1.0 / len(pos) for q in pos]
    ch = sum(wi * q["ch"] for wi, q in zip(w, pos))
    s2 = sum(wi * ((q["fw_ch"] / 2.3548200450309493) ** 2 + (q["ch"] - ch) ** 2) for wi, q in zip(w, pos))
    p = {"ch": ch, "fw_ch": 2.3548200450309493 * math.sqrt(s2),
         "net": sum(q["net"] for q in found), "unc": math.sqrt(sum(q["unc"] ** 2 for q in found)),
         "fw_kev": pos[0].get("fw_kev") if len(pos) == 1 else None,
         "e": pos[0].get("e") if len(pos) == 1 else None}
    e = _poly(coeffs, ch) if coeffs else p["e"]
    if coeffs:
        fw = abs(_dpoly(coeffs, p["ch"])) * p["fw_ch"]
    else:
        fw = p.get("fw_kev")
    rate = p["net"] / m.live_time if m.live_time else None
    rate_unc = p["unc"] / m.live_time if m.live_time else None
    ref = ref_date if ref_date is not None else track.effective_reference_date()
    dec = 1.0 if track.kind == "background" else decay_factor(line.half_life(), m.date, ref, m.real_time)
    rc = rate * dec if rate is not None else None
    eff = None
    act = track.reference_activity.get(line.nuclide)
    em = line.emission()
    if rc is not None and act and em:
        eff = rc / (act * em / 100.0)
    return LineResult(centroid_ch=p["ch"], centroid_kev=e, fwhm_kev=fw, net=p["net"], net_unc=p["unc"],
                      rate=rate, rate_unc=rate_unc, rate_corr=rc, decay=dec, efficiency=eff,
                      components=len(found))


@dataclass
class SeriesPoint:
    mid: str
    date: Optional[datetime]
    value: Optional[float]
    unc: Optional[float] = None
    excluded: bool = False
    baseline: bool = False
    z: Optional[float] = None
    flags: list[str] = field(default_factory=list)
    status: str = STATUS_OK


@dataclass
class Series:
    key: str                         # "rate@661.657" / "background"
    metric: str
    line: Optional[MonitoredLine]
    unit: str
    points: list[SeriesPoint] = field(default_factory=list)
    mean: Optional[float] = None
    sd: Optional[float] = None
    n_baseline: int = 0
    nominal: Optional[float] = None
    tol_lo: Optional[float] = None
    tol_hi: Optional[float] = None

    @property
    def title(self) -> str:
        return f"{self.metric} {self.line.label}" if self.line else self.metric

    def limits(self, k: float) -> tuple[Optional[float], Optional[float]]:
        if self.mean is None or not self.sd:
            return None, None
        return self.mean - k * self.sd, self.mean + k * self.sd

    def valid_points(self) -> list[SeriesPoint]:
        return [p for p in self.points if p.value is not None]

    def slope_per_year(self) -> Optional[float]:
        """Least-squares trend of the non-excluded points (unit per year)."""
        pts = [p for p in self.points if p.value is not None and not p.excluded and p.date]
        if len(pts) < 3:
            return None
        t0 = pts[0].date
        xs = [(p.date - t0).total_seconds() / (365.25 * _SECONDS_PER_DAY) for p in pts]
        ys = [p.value for p in pts]
        mx, my = sum(xs) / len(xs), sum(ys) / len(ys)
        sxx = sum((x - mx) ** 2 for x in xs)
        if sxx <= 0:
            return None
        return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sxx


@dataclass
class Violation:
    mid: str
    date: Optional[datetime]
    series: str
    rule: str
    value: Optional[float]
    message: str
    severity: str = STATUS_FAIL


@dataclass
class QCEvaluation:
    track: QCTrack
    series: list[Series]
    status: dict[str, str]                    # measurement id -> OK/WARN/FAIL/EXCL/N/A
    flags: dict[str, list[str]]               # measurement id -> ["rate@661.657:1-3s", ...]
    violations: list[Violation]
    line_results: dict[str, dict[str, Optional[LineResult]]]   # mid -> line key -> result
    reference_date: Optional[datetime] = None

    def get(self, metric: str, line=None) -> Optional[Series]:
        for s in self.series:
            if s.metric != metric:
                continue
            if line is None and s.line is None:
                return s
            if s.line is not None and line is not None:
                k = line.key if isinstance(line, MonitoredLine) else line
                if s.line.key == k or (isinstance(k, (int, float)) and abs(s.line.energy - k) < 0.05):
                    return s
        return None

    def counts(self) -> dict[str, int]:
        out: dict[str, int] = {}
        for v in self.status.values():
            out[v] = out.get(v, 0) + 1
        return out


def _baseline_ids(track: QCTrack, ms: list[QCMeasurement]) -> set[str]:
    use = [m for m in ms if not m.excluded]
    if track.baseline_start or track.baseline_end:
        return {m.id for m in use if m.date
                and (track.baseline_start is None or m.date >= track.baseline_start)
                and (track.baseline_end is None or m.date <= track.baseline_end)}
    return {m.id for m in use[:max(int(track.baseline_n), 0)]}


def _stats(vals: list[float]) -> tuple[Optional[float], Optional[float]]:
    if not vals:
        return None, None
    mean = sum(vals) / len(vals)
    if len(vals) < 2:
        return mean, None
    sd = math.sqrt(sum((v - mean) ** 2 for v in vals) / (len(vals) - 1))
    return mean, sd


def _apply_rules(s: Series, rules: Sequence[str]) -> None:
    pts = [p for p in s.points if p.value is not None and not p.excluded]
    if s.mean is not None and s.sd:
        for p in pts:
            p.z = (p.value - s.mean) / s.sd
        zs = [p.z for p in pts]
        rules = set(rules)
        for i, p in enumerate(pts):
            z = zs[i]
            if abs(z) > 3 and "1-3s" in rules:
                p.flags.append("1-3s")
            elif abs(z) > 2:
                p.flags.append("1-2s")
            if i >= 1:
                a = zs[i - 1]
                if "2-2s" in rules and ((a > 2 and z > 2) or (a < -2 and z < -2)):
                    p.flags.append("2-2s")
                if "R-4s" in rules and ((a > 2 and z < -2) or (a < -2 and z > 2)):
                    p.flags.append("R-4s")
            if i >= 3 and "4-1s" in rules:
                w = zs[i - 3:i + 1]
                if all(v > 1 for v in w) or all(v < -1 for v in w):
                    p.flags.append("4-1s")
            if i >= 9 and "10-x" in rules:
                w = zs[i - 9:i + 1]
                if all(v > 0 for v in w) or all(v < 0 for v in w):
                    p.flags.append("10-x")
    for p in pts:
        lo_bad = s.tol_lo is not None and p.value < s.tol_lo
        hi_bad = s.tol_hi is not None and p.value > s.tol_hi
        if lo_bad or hi_bad:
            p.flags.append("TOL")
    for p in s.points:
        if p.excluded:
            p.status = STATUS_EXCLUDED
        elif p.value is None:
            p.status = STATUS_NA
        elif any(f not in WARNING_RULES for f in p.flags):
            p.status = STATUS_FAIL
        elif p.flags:
            p.status = STATUS_WARN
        else:
            p.status = STATUS_OK


def _rule_text(rule: str) -> str:
    return {"1-2s": "beyond 2 sigma (warning)", "1-3s": "beyond 3 sigma",
            "2-2s": "2 consecutive beyond 2 sigma, same side",
            "R-4s": "consecutive points differ by more than 4 sigma",
            "4-1s": "4 consecutive beyond 1 sigma, same side",
            "10-x": "10 consecutive on one side of the mean",
            "TOL": "outside the fixed tolerance",
            "SPLIT": "line split into several peaks (peak doubling / zero or gain jump)"}.get(rule, rule)


def evaluate(track: QCTrack) -> QCEvaluation:
    ms = track.sorted_measurements()
    ref = track.effective_reference_date()
    base = _baseline_ids(track, ms)
    lr: dict[str, dict[str, Optional[LineResult]]] = {}
    for m in ms:
        coeffs = track.effective_coeffs(m)
        lr[m.id] = {l.key: line_result(track, m, l, ref, coeffs) for l in track.lines}

    series: list[Series] = []

    def build(key, metric, line, getter, unc_getter=None):
        s = Series(key=key, metric=metric, line=line, unit=METRIC_UNITS[metric])
        for m in ms:
            v = getter(m)
            s.points.append(SeriesPoint(mid=m.id, date=m.date, value=v,
                                        unc=unc_getter(m) if unc_getter else None,
                                        excluded=m.excluded, baseline=m.id in base))
        s.mean, s.sd = _stats([p.value for p in s.points if p.baseline and p.value is not None])
        s.n_baseline = sum(1 for p in s.points if p.baseline and p.value is not None)
        series.append(s)
        return s

    for l in track.lines:
        k = l.key
        s = build(f"centroid@{k}", "centroid", l,
                  lambda m, k=k: lr[m.id][k].centroid_kev if lr[m.id][k] else None)
        s.nominal = l.energy
        if track.centroid_tol_kev is not None:
            s.tol_lo, s.tol_hi = l.energy - track.centroid_tol_kev, l.energy + track.centroid_tol_kev
        s = build(f"fwhm@{k}", "fwhm", l, lambda m, k=k: lr[m.id][k].fwhm_kev if lr[m.id][k] else None)
        s.tol_hi = l.fwhm_max
        for pt in s.points:      # line found split into several peaks (peak doubling)
            r = lr[pt.mid][k]
            if r is not None and r.components > 1 and not pt.excluded:
                pt.flags.append("SPLIT")
        s = build(f"rate@{k}", "rate", l, lambda m, k=k: lr[m.id][k].rate_corr if lr[m.id][k] else None,
                  lambda m, k=k: (lr[m.id][k].rate_unc * lr[m.id][k].decay)
                  if lr[m.id][k] and lr[m.id][k].rate_unc is not None else None)
        if track.rate_tol_pct is not None and s.mean is not None:
            s.tol_lo = s.mean * (1 - track.rate_tol_pct / 100.0)
            s.tol_hi = s.mean * (1 + track.rate_tol_pct / 100.0)
    s = build("background", "background", None, lambda m: m.window_rate)
    if track.bkg_tol_pct is not None and s.mean is not None:
        s.tol_lo = s.mean * (1 - track.bkg_tol_pct / 100.0)
        s.tol_hi = s.mean * (1 + track.bkg_tol_pct / 100.0)

    violations: list[Violation] = []
    flags: dict[str, list[str]] = {m.id: [] for m in ms}
    for s in series:
        _apply_rules(s, track.rules)
        for p in s.points:
            for f in p.flags:
                flags[p.mid].append(f"{s.key}:{f}")
                violations.append(Violation(
                    mid=p.mid, date=p.date, series=s.key, rule=f, value=p.value,
                    message=f"{s.title}: {_rule_text(f)}",
                    severity=STATUS_WARN if f in WARNING_RULES else STATUS_FAIL))
    status: dict[str, str] = {}
    for m in ms:
        if m.excluded:
            status[m.id] = STATUS_EXCLUDED
            continue
        st = [s.points[i].status for s in series for i, p in enumerate(s.points) if p.mid == m.id]
        missing = track.lines and all(lr[m.id][l.key] is None for l in track.lines)
        if STATUS_FAIL in st:
            status[m.id] = STATUS_FAIL
        elif missing:
            status[m.id] = STATUS_NA
        elif STATUS_WARN in st:
            status[m.id] = STATUS_WARN
        else:
            status[m.id] = STATUS_OK
    violations.sort(key=lambda v: (v.date or datetime.max, v.series))
    return QCEvaluation(track=track, series=series, status=status, flags=flags, violations=violations,
                        line_results=lr, reference_date=ref)


@dataclass
class ViolationRun:
    series: str
    rule: str
    first: Optional[datetime]
    last: Optional[datetime]
    n: int
    first_mid: str
    message: str
    severity: str = STATUS_FAIL


def violation_runs(ev: QCEvaluation, include_warnings: bool = False) -> list[ViolationRun]:
    """Condensed violations: consecutive points of one series firing the same rule
    become one run (first date, last date, count)."""
    out: list[ViolationRun] = []
    for s in ev.series:
        pts = [p for p in s.points if p.value is not None and not p.excluded]
        open_runs: dict[str, ViolationRun] = {}
        for p in pts:
            fired = set(p.flags)
            for rule in list(open_runs):
                if rule not in fired:
                    out.append(open_runs.pop(rule))
            for rule in p.flags:
                if rule in WARNING_RULES and not include_warnings:
                    continue
                r = open_runs.get(rule)
                if r is None:
                    open_runs[rule] = ViolationRun(
                        series=s.key, rule=rule, first=p.date, last=p.date, n=1, first_mid=p.mid,
                        message=f"{s.title}: {_rule_text(rule)}",
                        severity=STATUS_WARN if rule in WARNING_RULES else STATUS_FAIL)
                else:
                    r.last, r.n = p.date, r.n + 1
        out.extend(open_runs.values())
    out.sort(key=lambda r: (r.first or datetime.max, r.series, r.rule))
    return out


# ------------------------------------------------------------------ Excel export
def export_track_xlsx(track: QCTrack, path, ev: Optional[QCEvaluation] = None) -> Path:
    """Workbook with sheets Track / Measurements / Statistics / Violations."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    ev = ev or evaluate(track)
    path = Path(path)
    wb = Workbook()
    bold = Font(bold=True)
    ws = wb.active
    ws.title = "Track"
    info = [("Name", track.name), ("Detector", track.detector), ("Detector type", track.detector_type),
            ("Check source", track.source), ("Kind", track.kind),
            ("Monitored lines", ", ".join(l.label for l in track.lines)),
            ("Reference date", _iso(ev.reference_date) or ""),
            ("Reference activity (Bq)", ", ".join(f"{k}: {v:g}" for k, v in track.reference_activity.items())),
            ("Centroid mode", track.centroid_mode),
            ("Background window (keV)", f"{track.bkg_window[0] or 'start'} - {track.bkg_window[1] or 'end'}"),
            ("Baseline", (f"{_iso(track.baseline_start) or ''} .. {_iso(track.baseline_end) or ''}"
                          if track.baseline_start or track.baseline_end else f"first {track.baseline_n} points")),
            ("Westgard rules", ", ".join(track.rules)),
            ("Centroid tolerance (keV)", track.centroid_tol_kev),
            ("Rate tolerance (%)", track.rate_tol_pct), ("Background tolerance (%)", track.bkg_tol_pct),
            ("FWHM limits (keV)", ", ".join(f"{l.energy:.1f}: {l.fwhm_max}" for l in track.lines if l.fwhm_max)),
            ("Measurements", len(track.measurements)), ("Notes", track.notes),
            ("Exported", _iso(datetime.now()))]
    for k, v in info:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = bold
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70

    ws = wb.create_sheet("Measurements")
    head = ["Date", "File", "Sample ID", "Live (s)", "Real (s)", "Dead (%)", "Background (cps)"]
    for l in track.lines:
        n = f"{l.energy:.1f}"
        head += [f"Centroid {n} (ch)", f"Centroid {n} (keV)", f"FWHM {n} (keV)", f"Rate {n} (cps)",
                 f"Rate {n} decay-corr (cps)", f"Efficiency {n}"]
    head += ["Status", "Rules fired", "Excluded", "Note", "Path"]
    ws.append(head)
    for c in ws[1]:
        c.font = bold
    fills = {STATUS_FAIL: PatternFill("solid", fgColor="F8D7DA"), STATUS_WARN: PatternFill("solid", fgColor="FFF3CD")}
    for m in track.sorted_measurements():
        row = [m.date, m.file_name, m.sample_id, m.live_time, m.real_time,
               round(m.dead_time_pct, 3) if m.dead_time_pct is not None else None, m.window_rate]
        for l in track.lines:
            r = ev.line_results.get(m.id, {}).get(l.key)
            row += ([r.centroid_ch, r.centroid_kev, r.fwhm_kev, r.rate, r.rate_corr, r.efficiency]
                    if r else [None] * 6)
        st = ev.status.get(m.id, "")
        row += [st, "; ".join(ev.flags.get(m.id, [])), "yes" if m.excluded else "", m.note, m.source_path]
        ws.append(row)
        if st in fills:
            ws.cell(ws.max_row, len(head) - 4).fill = fills[st]
    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 22
    ws.freeze_panes = "C2"

    ws = wb.create_sheet("Statistics")
    ws.append(["Series", "Unit", "N baseline", "Mean", "SD", "CV (%)", "-3 SD", "-2 SD", "+2 SD", "+3 SD",
               "Tol. low", "Tol. high", "Trend (/year)", "N points", "N fail", "N warn"])
    for c in ws[1]:
        c.font = bold
    for s in ev.series:
        lo2, hi2 = s.limits(2)
        lo3, hi3 = s.limits(3)
        cv = 100 * s.sd / abs(s.mean) if s.sd and s.mean else None
        ws.append([s.title, s.unit, s.n_baseline, s.mean, s.sd, cv, lo3, lo2, hi2, hi3, s.tol_lo, s.tol_hi,
                   s.slope_per_year(), len(s.valid_points()),
                   sum(1 for p in s.points if p.status == STATUS_FAIL),
                   sum(1 for p in s.points if p.status == STATUS_WARN)])
    ws.column_dimensions["A"].width = 32

    ws = wb.create_sheet("Violations")
    ws.append(["Date", "File", "Series", "Rule", "Severity", "Value", "Message"])
    for c in ws[1]:
        c.font = bold
    names = {m.id: m.file_name for m in track.measurements}
    for v in ev.violations:
        ws.append([v.date, names.get(v.mid, ""), v.series, v.rule, v.severity, v.value, v.message])
    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["G"].width = 60
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
