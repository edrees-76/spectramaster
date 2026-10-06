"""Archive catalogue (index) and rescue (verified bulk conversion).

Both functions are strictly read-only on the archive: sources are only opened
for reading and the output location may not lie inside the archive root.

index_archive
    Walks a folder, reads every recognised spectrum and writes an Excel catalogue
    (sheets: Catalog, Summary, Duplicates, Mirrored folders, Errors).

rescue_archive
    Converts every spectrum into one or more modern formats, preserving the
    folder tree (``OUT/<format key>/<relative folder>/<name>``), reads every output
    back and compares it with the source (counts exact, live/real time, start
    time, energy at several channels, metadata), writes a verification report
    (xlsx) and a manifest (CSV with SHA-256 of sources and outputs).  The manifest
    is appended after every file, so an interrupted run can be resumed
    (``resume=True``): entries already verified whose source hash and output hash
    are unchanged are skipped.
"""
from __future__ import annotations

import csv
import hashlib
import os
import struct
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable, Optional, Sequence

from ..formats import FormatSpec, detect_format, get_format, load_all, readable_formats
from ..model import Spectrum

__all__ = [
    "STRONG_SPECTRUM_EXTENSIONS", "find_spectrum_files", "sha256_file", "spectrum_fingerprint",
    "CatalogEntry", "IndexResult", "scan_file", "index_archive",
    "Verification", "verify_output", "RescueEntry", "RescueResult", "rescue_archive",
]

ProgressCB = Callable[[int, int, Path], Optional[bool]]
"""progress(done, total, current_path) -> False to cancel."""

# Extensions that are spectra and nothing else: a file with one of these that no
# reader recognises is reported as an error instead of being silently skipped.
STRONG_SPECTRUM_EXTENSIONS = {".s0", ".chn", ".cnf", ".spe", ".n42", ".spc", ".mca", ".iec",
                              ".pcf", ".tka", ".apt", ".spm", ".spt", ".ans", ".de1", ".sct"}


# ------------------------------------------------------------------ helpers
def _norm(p: Path) -> str:
    return os.path.normcase(os.path.abspath(str(p)))


def _is_inside(child: Path, parent: Path) -> bool:
    c, p = _norm(child), _norm(parent)
    return c == p or c.startswith(p.rstrip("\\/") + os.sep)


def _guard_output(root: Path, out: Path, what: str) -> None:
    if _is_inside(out, root):
        raise ValueError(f"{what} {out} lies inside the archive {root}: the archive is read-only, "
                         "choose a location outside it")


def sha256_file(path: Path, chunk: int = 1 << 20) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        while True:
            b = fh.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def spectrum_fingerprint(s: Spectrum) -> str:
    """Hash of the measurement content (counts, live/real time, start time),
    independent of the file format / byte layout.  Used to find the same spectrum
    stored in different files."""
    h = hashlib.sha256()
    h.update(struct.pack(f"<{len(s.counts)}d", *[float(c) for c in s.counts]))
    h.update(repr((round(s.live_time or 0.0, 3), round(s.real_time or 0.0, 3),
                   s.start_time.isoformat(timespec="seconds") if s.start_time else "")).encode())
    return h.hexdigest()


def find_spectrum_files(root: Path, recursive: bool = True, all_files: bool = False
                        ) -> list[Path]:
    """Candidate files under `root` (sorted).  Unless `all_files`, only files whose
    extension belongs to some readable format are returned (skips .doc, .xls,
    .pdf ... cheaply; content sniffing happens later)."""
    load_all()
    exts = set()
    for f in readable_formats():
        exts.update(f.extensions)
    exts |= STRONG_SPECTRUM_EXTENSIONS
    root = Path(root)
    out = []
    if recursive:
        walker = os.walk(root)
    else:
        walker = [(str(root), [], [n for n in os.listdir(root) if (root / n).is_file()])]
    for dirpath, dirnames, filenames in walker:
        dirnames.sort(key=str.lower)
        for n in filenames:
            if all_files or os.path.splitext(n)[1].lower() in exts:
                out.append(Path(dirpath) / n)
    out.sort(key=lambda p: str(p).lower())
    return out


def _recognise(path: Path) -> Optional[FormatSpec]:
    return detect_format(path)


# ------------------------------------------------------------------- index
@dataclass
class CatalogEntry:
    path: Path
    rel_path: str
    folder: str
    file_name: str
    file_size: int
    sha256: str
    format_key: str = ""
    format_name: str = ""
    spectrum_index: int = 1
    spectra_in_file: int = 1
    sample_id: str = ""
    description: str = ""
    detector: str = ""
    start_time: Optional[datetime] = None
    live_time: Optional[float] = None
    real_time: Optional[float] = None
    dead_time_percent: Optional[float] = None
    channels: int = 0
    total_counts: Optional[float] = None
    calibrated: bool = False
    energy_coeffs: list[float] = field(default_factory=list)
    energy_range: str = ""
    fwhm_coeffs: list[float] = field(default_factory=list)
    fingerprint: str = ""
    status: str = "OK"
    error: bool = False

    @property
    def year(self) -> Optional[int]:
        return self.start_time.year if self.start_time else None


def _energies(s: Spectrum):
    """numpy energies of the stored channels (fast), or None."""
    if not s.energy_coeffs or not s.counts:
        return None
    import numpy as np
    from numpy.polynomial import polynomial as P
    return P.polyval(np.arange(s.first_channel, s.first_channel + s.num_channels, dtype=float),
                     s.energy_coeffs)


def _checks(s: Spectrum, e=None) -> list[str]:
    import numpy as np
    w = []
    if not s.counts:
        w.append("no channels")
    elif sum(s.counts) <= 0:
        w.append("empty spectrum (0 counts)")
    if s.counts and min(s.counts) < 0:
        w.append("negative counts")
    if not s.energy_coeffs:
        w.append("uncalibrated")
    elif e is not None and len(e) > 1 and np.any(np.diff(e) <= 0):
        w.append("non-monotonic energy calibration")
    if s.live_time is None or s.live_time <= 0:
        w.append("no live time")
    if s.real_time is None or s.real_time <= 0:
        w.append("no real time")
    if s.live_time and s.real_time and s.live_time > s.real_time * 1.0001:
        w.append("live time > real time")
    if s.start_time is None:
        w.append("no start time")
    elif not (1980 <= s.start_time.year <= datetime.now().year + 1):
        w.append(f"implausible start date {s.start_time.date()}")
    return w


def scan_file(path: Path, root: Optional[Path] = None, fmt: Optional[FormatSpec] = None,
              strong_only_errors: bool = True) -> list[CatalogEntry]:
    """Catalogue one file: one entry per spectrum, or one error entry.  Returns []
    if the file is not a recognised spectrum and its extension is not a
    spectrum-only one (see STRONG_SPECTRUM_EXTENSIONS)."""
    load_all()
    path = Path(path)
    root = Path(root) if root else path.parent
    try:
        rel = os.path.relpath(path, root)
    except ValueError:
        rel = str(path)
    folder = os.path.dirname(rel)
    try:
        size = path.stat().st_size
    except OSError as exc:
        return [CatalogEntry(path, rel, folder, path.name, 0, "", status=f"error: {exc}", error=True)]
    base = dict(path=path, rel_path=rel, folder=folder, file_name=path.name, file_size=size)
    if fmt is None:
        fmt = _recognise(path)
    if fmt is None:
        if path.suffix.lower() in STRONG_SPECTRUM_EXTENSIONS or not strong_only_errors:
            return [CatalogEntry(**base, sha256=_safe_hash(path),
                                 status="error: unrecognised or corrupt spectrum file", error=True)]
        return []
    digest = _safe_hash(path)
    try:
        spectra = fmt.reader(path)
        if not spectra:
            raise ValueError("no spectrum in file")
    except Exception as exc:
        return [CatalogEntry(**base, sha256=digest, format_key=fmt.key, format_name=fmt.name,
                             status=f"error: {type(exc).__name__}: {exc}", error=True)]
    out = []
    for i, s in enumerate(spectra, start=1):
        e = _energies(s)
        warns = _checks(s, e)
        if len(spectra) > 1:
            warns.insert(0, f"spectrum {i} of {len(spectra)}")
        out.append(CatalogEntry(
            **base, sha256=digest, format_key=fmt.key, format_name=fmt.name, spectrum_index=i,
            spectra_in_file=len(spectra), sample_id=s.sample_id, description=s.description,
            detector=s.detector, start_time=s.start_time, live_time=s.live_time, real_time=s.real_time,
            dead_time_percent=s.dead_time_percent, channels=s.num_channels, total_counts=s.total_counts,
            calibrated=bool(s.energy_coeffs), energy_coeffs=list(s.energy_coeffs),
            energy_range=f"{e[0]:.2f} - {e[-1]:.2f}" if e is not None else "", fwhm_coeffs=list(s.fwhm_coeffs),
            fingerprint=spectrum_fingerprint(s),
            status="OK" if not warns else "warning: " + "; ".join(warns)))
    return out


def _safe_hash(path: Path) -> str:
    try:
        return sha256_file(path)
    except OSError:
        return ""


@dataclass
class DuplicateGroup:
    kind: str                       # "identical file" / "same spectrum data"
    key: str
    entries: list[CatalogEntry]


@dataclass
class IndexResult:
    root: Path
    entries: list[CatalogEntry]
    duplicates: list[DuplicateGroup]
    mirrored_folders: list[tuple[str, str, int]]
    files_scanned: int
    files_skipped: int
    elapsed: float
    output: Optional[Path] = None
    cancelled: bool = False

    @property
    def errors(self) -> list[CatalogEntry]:
        return [e for e in self.entries if e.error]

    @property
    def duplicate_files(self) -> int:
        """Number of redundant copies (files beyond the first of each identical group)."""
        return sum(len({x.rel_path for x in g.entries}) - 1 for g in self.duplicates
                   if g.kind == "identical file")


def find_duplicates(entries: Sequence[CatalogEntry]) -> tuple[list[DuplicateGroup], list[tuple[str, str, int]]]:
    """Group identical files (SHA-256) and identical spectrum content stored in
    byte-different files; also count file pairs shared between folder pairs
    (mirrored trees)."""
    by_hash: dict[str, list[CatalogEntry]] = defaultdict(list)
    for e in entries:
        if e.sha256 and e.spectrum_index == 1:
            by_hash[e.sha256].append(e)
    groups = [DuplicateGroup("identical file", h, v) for h, v in by_hash.items() if len(v) > 1]
    by_fp: dict[str, list[CatalogEntry]] = defaultdict(list)
    for e in entries:
        if e.fingerprint and not e.error:
            by_fp[e.fingerprint].append(e)
    for fp, v in by_fp.items():
        if len({x.sha256 for x in v}) > 1:
            groups.append(DuplicateGroup("same spectrum data", fp, v))
    groups.sort(key=lambda g: (g.kind, g.entries[0].rel_path.lower()))
    pairs: Counter = Counter()
    for g in groups:
        if g.kind != "identical file":
            continue
        folders = sorted({e.folder for e in g.entries})
        for i in range(len(folders)):
            for j in range(i + 1, len(folders)):
                pairs[(folders[i], folders[j])] += 1
    mirrored = [(a, b, n) for (a, b), n in pairs.most_common()]
    return groups, mirrored


def index_archive(root, out_xlsx=None, recursive: bool = True, progress: Optional[ProgressCB] = None,
                  all_files: bool = False, files: Optional[Iterable[Path]] = None) -> IndexResult:
    """Catalogue every recognised spectrum under `root` and (if `out_xlsx`) write
    the Excel catalogue.  `files` restricts the scan to the given paths (still
    reported relative to `root`).  Returns an IndexResult (also usable by a GUI
    table without writing Excel)."""
    load_all()
    root = Path(root)
    if out_xlsx is not None:
        _guard_output(root, Path(out_xlsx), "catalogue")
    t0 = time.perf_counter()
    cand = [Path(p) for p in files] if files is not None else find_spectrum_files(root, recursive, all_files)
    entries: list[CatalogEntry] = []
    skipped = 0
    cancelled = False
    for i, p in enumerate(cand, start=1):
        rows = scan_file(p, root)
        if not rows:
            skipped += 1
        entries.extend(rows)
        if progress is not None and progress(i, len(cand), p) is False:
            cancelled = True
            break
    dups, mirrored = find_duplicates(entries)
    res = IndexResult(root, entries, dups, mirrored, len(cand), skipped, time.perf_counter() - t0,
                      cancelled=cancelled)
    if out_xlsx is not None:
        write_catalog_xlsx(res, Path(out_xlsx))
        res.output = Path(out_xlsx)
    return res


# ------------------------------------------------------------ excel output
def _xl_safe(v):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("?", v)
        return v[:32000]
    return v


def _sheet(wb, title, header, rows, widths=None, first=False):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    fill = PatternFill("solid", fgColor="DDE7F3")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = fill
        c.alignment = Alignment(vertical="center", wrap_text=False)
    for r in rows:
        ws.append([_xl_safe(v) for v in r])
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(header, start=1):
        w = (widths or {}).get(h)
        if w is None:
            sample = [len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if r[i - 1] is not None]
            w = min(max(sample) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def _fmt_coeffs(c: Sequence[float]) -> str:
    return ", ".join(f"{x:.9g}" for x in c)


CATALOG_COLUMNS = ["Path", "Folder", "File name", "Format", "Spectrum #", "Sample ID / title",
                   "Description", "Detector", "Start date/time", "Year", "Live time (s)", "Real time (s)",
                   "Dead time (%)", "Channels", "Total counts", "Calibrated", "Energy coefficients",
                   "Energy range (keV)", "FWHM coefficients", "File size (bytes)", "SHA-256", "Status"]


def write_catalog_xlsx(res: IndexResult, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    rows = []
    for e in res.entries:
        rows.append([str(e.path), e.folder, e.file_name, e.format_key, e.spectrum_index, e.sample_id,
                     e.description.replace("\r", " ").replace("\n", " | "), e.detector, e.start_time,
                     e.year, e.live_time, e.real_time,
                     round(e.dead_time_percent, 3) if e.dead_time_percent is not None else None,
                     e.channels or None, e.total_counts, "y" if e.calibrated else ("n" if not e.error else ""),
                     _fmt_coeffs(e.energy_coeffs), e.energy_range, _fmt_coeffs(e.fwhm_coeffs),
                     e.file_size, e.sha256, e.status])
    ws = _sheet(wb, "Catalog", CATALOG_COLUMNS, rows, first=True,
                widths={"Path": 60, "Folder": 45, "File name": 34, "Description": 40, "SHA-256": 20,
                        "Status": 45, "Start date/time": 19, "Energy coefficients": 40,
                        "FWHM coefficients": 30, "Sample ID / title": 26})
    for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
        for c in row:
            c.number_format = "yyyy-mm-dd hh:mm:ss"
    red = PatternFill("solid", fgColor="F8D7DA")
    amber = PatternFill("solid", fgColor="FFF3CD")
    status_col = len(CATALOG_COLUMNS)
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, status_col)
        if isinstance(c.value, str) and c.value.startswith("error"):
            c.fill = red
        elif isinstance(c.value, str) and c.value.startswith("warning"):
            c.fill = amber

    # summary
    ok = [e for e in res.entries if not e.error]
    summ = wb.create_sheet("Summary")
    bold = Font(bold=True)

    def block(title, counter: Counter, key_name):
        summ.append([])
        summ.append([title])
        summ.cell(summ.max_row, 1).font = bold
        summ.append([key_name, "Spectra"])
        for k, v in sorted(counter.items(), key=lambda kv: (str(kv[0]))):
            summ.append([_xl_safe(k if k not in (None, "") else "(unknown)"), v])

    summ.append(["Archive root", str(res.root)])
    summ.append(["Indexed at", datetime.now().replace(microsecond=0)])
    summ.append(["Candidate files scanned", res.files_scanned])
    summ.append(["Files not recognised as spectra (skipped)", res.files_skipped])
    summ.append(["Spectra catalogued", len(ok)])
    summ.append(["Files with errors", len(res.errors)])
    summ.append(["Spectra with warnings", sum(1 for e in ok if e.status != "OK")])
    summ.append(["Uncalibrated spectra", sum(1 for e in ok if not e.calibrated)])
    summ.append(["Identical-file duplicate groups", sum(1 for g in res.duplicates if g.kind == "identical file")])
    summ.append(["Redundant identical copies", res.duplicate_files])
    summ.append(["Same-data groups (different bytes)",
                 sum(1 for g in res.duplicates if g.kind == "same spectrum data")])
    summ.append(["Elapsed (s)", round(res.elapsed, 1)])
    if res.cancelled:
        summ.append(["NOTE", "scan was cancelled: catalogue is incomplete"])
    for r in range(1, summ.max_row + 1):
        summ.cell(r, 1).font = bold
    block("Spectra per year", Counter(e.year for e in ok), "Year")
    block("Spectra per detector", Counter(e.detector for e in ok), "Detector")
    block("Spectra per format", Counter(e.format_key for e in ok), "Format")
    block("Status", Counter(e.status.split(":")[0] for e in res.entries), "Status")
    warn_kinds: Counter = Counter()
    for e in ok:
        if e.status.startswith("warning: "):
            for w in e.status[9:].split("; "):
                warn_kinds[w.split(" (")[0] if not w.startswith("implausible") else "implausible start date"] += 1
    block("Warnings", warn_kinds, "Warning")
    summ.column_dimensions["A"].width = 48
    summ.column_dimensions["B"].width = 60

    # duplicates
    drows = []
    for gi, g in enumerate(res.duplicates, start=1):
        for e in g.entries:
            drows.append([gi, g.kind, len(g.entries), str(e.path), e.folder, e.file_name, e.format_key,
                          e.start_time, e.total_counts, e.sha256 if g.kind == "identical file" else g.key])
    _sheet(wb, "Duplicates", ["Group", "Kind", "Copies", "Path", "Folder", "File name", "Format",
                              "Start date/time", "Total counts", "SHA-256 / content hash"], drows,
           widths={"Path": 60, "Folder": 45, "File name": 34, "SHA-256 / content hash": 20, "Start date/time": 19})
    _sheet(wb, "Mirrored folders", ["Folder A", "Folder B", "Identical files shared"],
           [list(x) for x in res.mirrored_folders], widths={"Folder A": 70, "Folder B": 70})
    _sheet(wb, "Errors", ["Path", "Format", "File size (bytes)", "Error"],
           [[str(e.path), e.format_key, e.file_size, e.status] for e in res.errors],
           widths={"Path": 80, "Error": 80})
    wb.save(out)
    return out


# ------------------------------------------------------------------ rescue
@dataclass
class Verification:
    ok: bool
    problems: list[str] = field(default_factory=list)       # failures
    warnings: list[str] = field(default_factory=list)
    counts_equal: bool = False
    live_dev: Optional[float] = None       # s
    real_dev: Optional[float] = None       # s
    start_dev: Optional[float] = None      # s
    energy_dev: Optional[float] = None     # keV, max over the probe channels

    @property
    def status(self) -> str:
        if not self.ok:
            return "FAILED"
        return "VERIFIED (warnings)" if self.warnings else "VERIFIED"


def verify_output(src: Spectrum, out: Spectrum, time_tol: float = 0.02, time_rel_tol: float = 1e-6,
                  start_tol: float = 1.0, energy_tol: float = 0.05, energy_fail: float = 1.0,
                  probe_fractions: Sequence[float] = (0.0, 0.1, 0.25, 0.5, 0.75, 0.9, 1.0)) -> Verification:
    """Compare a re-read output spectrum with its source.

    Failures: channel count / counts differ (exact), live/real time beyond
    max(time_tol, time_rel_tol*t) seconds, start time beyond `start_tol` s, a
    calibration lost or an energy deviation > `energy_fail` keV (default 1).  Warnings: energy deviation
    above `energy_tol` keV (typically a refit to fewer terms), changed
    sample ID / description / detector text.
    """
    v = Verification(ok=True)
    if out.num_channels != src.num_channels:
        v.problems.append(f"channels {src.num_channels} -> {out.num_channels}")
    else:
        diff = [i for i, (a, b) in enumerate(zip(src.counts, out.counts)) if a != b]
        v.counts_equal = not diff
        if diff:
            v.problems.append(f"counts differ in {len(diff)} channel(s), first at index {diff[0]} "
                              f"({src.counts[diff[0]]} -> {out.counts[diff[0]]})")
    for name, a, b in (("live time", src.live_time, out.live_time), ("real time", src.real_time, out.real_time)):
        if a is None:
            continue
        if b is None:
            v.problems.append(f"{name} lost")
            continue
        d = abs(a - b)
        if name == "live time":
            v.live_dev = d
        else:
            v.real_dev = d
        if d > max(time_tol, time_rel_tol * abs(a)):
            v.problems.append(f"{name} {a} -> {b} (diff {d:.4g} s)")
    if src.start_time is not None:
        if out.start_time is None:
            v.problems.append("start time lost")
        else:
            v.start_dev = abs((out.start_time - src.start_time).total_seconds())
            if v.start_dev > start_tol:
                v.problems.append(f"start time {src.start_time} -> {out.start_time}")
    if src.energy_coeffs:
        if not out.energy_coeffs:
            v.problems.append("energy calibration lost")
        else:
            n = src.num_channels
            chans = sorted({src.first_channel + int(round(f * (n - 1))) for f in probe_fractions})
            dev = max(abs(src.energy_of(c) - out.energy_of(c - src.first_channel + out.first_channel))
                      for c in chans)
            v.energy_dev = dev
            if dev > energy_fail:
                v.problems.append(f"energy deviates by {dev:.3f} keV")
            elif dev > energy_tol:
                v.warnings.append(f"energy deviates by up to {dev:.3f} keV (calibration refitted/rounded)")
    for name in ("sample_id", "description", "detector"):
        a, b = (getattr(src, name) or "").strip(), (getattr(out, name) or "").strip()
        if a and a != b:
            v.warnings.append(f"{name.replace('_', ' ')} changed ({a[:30]!r} -> {b[:30]!r})")
    v.ok = not v.problems
    return v


MANIFEST_FIELDS = ["source_rel_path", "spectrum_index", "format", "source_sha256", "source_size",
                   "output_rel_path", "output_sha256", "output_size", "status", "counts_equal",
                   "live_dev_s", "real_dev_s", "start_dev_s", "energy_dev_keV", "details", "processed_at"]


@dataclass
class RescueEntry:
    source: Path
    source_rel: str
    spectrum_index: int
    format_key: str
    source_sha256: str = ""
    source_size: int = 0
    output: Optional[Path] = None
    output_rel: str = ""
    output_sha256: str = ""
    output_size: int = 0
    status: str = ""                 # VERIFIED / VERIFIED (warnings) / FAILED / ERROR / WRITTEN / DUPLICATE
    details: str = ""
    verification: Optional[Verification] = None
    skipped: bool = False            # resumed from the manifest, not redone
    processed_at: str = ""
    row: dict = field(default_factory=dict)

    @property
    def key(self) -> tuple:
        return (self.source_rel.lower(), str(self.spectrum_index), self.format_key)

    def to_row(self) -> dict:
        if self.skipped and self.row:
            return dict(self.row)
        v = self.verification
        f = lambda x: "" if x is None else f"{x:.6g}"
        return {"source_rel_path": self.source_rel, "spectrum_index": self.spectrum_index,
                "format": self.format_key, "source_sha256": self.source_sha256,
                "source_size": self.source_size, "output_rel_path": self.output_rel,
                "output_sha256": self.output_sha256, "output_size": self.output_size,
                "status": self.status, "counts_equal": "" if v is None else ("yes" if v.counts_equal else "no"),
                "live_dev_s": f(v.live_dev if v else None), "real_dev_s": f(v.real_dev if v else None),
                "start_dev_s": f(v.start_dev if v else None), "energy_dev_keV": f(v.energy_dev if v else None),
                "details": self.details, "processed_at": self.processed_at}


@dataclass
class RescueResult:
    root: Path
    out_root: Path
    entries: list[RescueEntry]
    files_total: int
    elapsed: float
    manifest: Path
    report: Optional[Path]
    cancelled: bool = False

    def counts(self) -> Counter:
        return Counter(("skipped " if e.skipped else "") + e.status for e in self.entries)


def _read_manifest(path: Path) -> dict[tuple, dict]:
    rows: dict[tuple, dict] = {}
    if not path.exists():
        return rows
    with open(path, newline="", encoding="utf-8-sig") as fh:
        for r in csv.DictReader(fh):
            if not r.get("source_rel_path"):
                continue
            rows[(r["source_rel_path"].lower(), r["spectrum_index"], r["format"])] = r
    return rows


def rescue_archive(root, out_root, formats: Sequence[str] = ("n42_2012", "ci_cnf"), keep_tree: bool = True,
                   verify: bool = True, resume: bool = False, recursive: bool = True,
                   skip_duplicates: bool = False, progress: Optional[ProgressCB] = None,
                   files: Optional[Iterable[Path]] = None, report_xlsx=None, manifest_csv=None,
                   energy_tol: float = 0.05, time_tol: float = 0.02, all_files: bool = False,
                   energy_fail: float = 1.0) -> RescueResult:
    """Convert every spectrum under `root` into each of `formats` below
    ``out_root/<format key>/`` (keeping the relative folder tree if `keep_tree`),
    verify each output by reading it back, and write ``manifest.csv`` and
    ``verification_report.xlsx`` into `out_root` (or the given paths).

    resume=True: an existing manifest is loaded; entries with status VERIFIED*
    whose source SHA-256 is unchanged and whose output file still has the
    recorded SHA-256 are skipped.  The manifest is appended after each file, so
    an interrupted run loses at most the file in progress.

    skip_duplicates=True: a source byte-identical to one already processed in
    this run is not converted again (status DUPLICATE, details = first copy).
    """
    from ..convert import output_name, write_spectrum

    load_all()
    root, out_root = Path(root), Path(out_root)
    _guard_output(root, out_root, "output folder")
    fmts = [get_format(f) if not isinstance(f, FormatSpec) else f for f in formats]
    for f in fmts:
        if not f.can_write or not f.can_read:
            raise ValueError(f"{f.key}: rescue formats must be readable and writable (for verification)")
    out_root.mkdir(parents=True, exist_ok=True)
    manifest = Path(manifest_csv) if manifest_csv else out_root / "manifest.csv"
    report = Path(report_xlsx) if report_xlsx else out_root / "verification_report.xlsx"
    previous = _read_manifest(manifest) if resume else {}
    prev_by_rel: dict[str, list[dict]] = defaultdict(list)
    for k, r in previous.items():
        prev_by_rel[k[0]].append(r)
    if not resume and manifest.exists():
        manifest.unlink()
    new_file = not manifest.exists()
    mfh = open(manifest, "a", newline="", encoding="utf-8-sig" if new_file else "utf-8")
    writer = csv.DictWriter(mfh, fieldnames=MANIFEST_FIELDS)
    if new_file:
        writer.writeheader()
    t0 = time.perf_counter()
    cand = [Path(p) for p in files] if files is not None else find_spectrum_files(root, recursive, all_files)
    entries: list[RescueEntry] = []
    seen_hash: dict[str, str] = {}
    cancelled = False
    out_hash_cache: dict[str, str] = {}

    def rel_of(p: Path) -> str:
        try:
            return os.path.relpath(p, root)
        except ValueError:
            return p.name

    try:
        for n, src in enumerate(cand, start=1):
            rel = rel_of(src)
            stamp = datetime.now().replace(microsecond=0).isoformat(sep=" ")
            file_entries: list[RescueEntry] = []
            fmt_in = _recognise(src)
            if fmt_in is None and src.suffix.lower() not in STRONG_SPECTRUM_EXTENSIONS:
                if progress is not None and progress(n, len(cand), src) is False:
                    cancelled = True
                    break
                continue            # not a spectrum (report, table ...): ignore silently
            try:
                digest = sha256_file(src)
                size = src.stat().st_size
            except OSError as exc:
                file_entries.append(RescueEntry(src, rel, 1, "", status="ERROR", details=str(exc), processed_at=stamp))
                digest, size = "", 0
            if digest and fmt_in is None:
                file_entries.append(RescueEntry(src, rel, 1, "", digest, size, status="ERROR",
                                                details="unrecognised or corrupt spectrum file",
                                                processed_at=stamp))
            elif digest:
                # resume: everything for this file already verified & intact?
                prev = prev_by_rel.get(rel.lower(), [])
                if prev and _all_intact(prev, fmts, digest, out_root, out_hash_cache):
                    for r in prev:
                        e = RescueEntry(src, rel, int(r["spectrum_index"]), r["format"], digest, size,
                                        out_root / r["output_rel_path"], r["output_rel_path"],
                                        r["output_sha256"], int(r["output_size"] or 0), r["status"],
                                        r["details"], skipped=True, processed_at=r["processed_at"], row=r)
                        entries.append(e)
                    if progress is not None and progress(n, len(cand), src) is False:
                        cancelled = True
                        break
                    continue
                if skip_duplicates and digest in seen_hash:
                    file_entries.append(RescueEntry(src, rel, 1, "*", digest, size, status="DUPLICATE",
                                                    details=f"identical to {seen_hash[digest]}", processed_at=stamp))
                else:
                    seen_hash.setdefault(digest, rel)
                    file_entries.extend(_rescue_one(src, rel, digest, size, fmt_in, fmts, out_root, keep_tree,
                                                    verify, (energy_tol, energy_fail), time_tol, stamp, output_name,
                                                    write_spectrum))
            for e in file_entries:
                writer.writerow(e.to_row())
            mfh.flush()
            entries.extend(file_entries)
            if progress is not None and progress(n, len(cand), src) is False:
                cancelled = True
                break
    finally:
        mfh.close()
    # compact the journal: one row per (source, spectrum, format), latest wins
    final = _read_manifest(manifest)
    for e in entries:
        if e.skipped:
            final[e.key] = e.to_row()
    with open(manifest, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=MANIFEST_FIELDS)
        w.writeheader()
        for k in sorted(final):
            w.writerow({f: final[k].get(f, "") for f in MANIFEST_FIELDS})
    res = RescueResult(root, out_root, entries, len(cand), time.perf_counter() - t0, manifest, None, cancelled)
    if report is not None:
        _guard_output(root, report, "report")
        write_rescue_report(res, report, final)
        res.report = report
    return res


def _all_intact(prev_rows, fmts, digest, out_root, cache) -> bool:
    keys = {f.key for f in fmts}
    have = {r["format"] for r in prev_rows}
    if not keys <= have:
        return False
    for r in prev_rows:
        if r["format"] not in keys:
            continue
        if not r["status"].startswith("VERIFIED") or r["source_sha256"] != digest:
            return False
        out = out_root / r["output_rel_path"]
        if not out.is_file():
            return False
        k = _norm(out)
        if k not in cache:
            cache[k] = sha256_file(out)
        if cache[k] != r["output_sha256"]:
            return False
    return True


def _rescue_one(src, rel, digest, size, fmt_in, fmts, out_root, keep_tree, verify, energy_tols, time_tol,
                stamp, output_name, write_spectrum) -> list[RescueEntry]:
    out: list[RescueEntry] = []
    try:
        spectra = fmt_in.reader(src)
        if not spectra:
            raise ValueError("no spectrum in file")
    except Exception as exc:
        return [RescueEntry(src, rel, 1, "", digest, size, status="ERROR",
                            details=f"read failed ({fmt_in.key}): {type(exc).__name__}: {exc}", processed_at=stamp)]
    sub = Path(os.path.dirname(rel)) if keep_tree else Path()
    for i, s in enumerate(spectra, start=1):
        s.source_path, s.source_format = str(src), fmt_in.key
        for f in fmts:
            target = out_root / f.key / sub / output_name(src, f, i if len(spectra) > 1 else None)
            e = RescueEntry(src, rel, i, f.key, digest, size, target,
                            os.path.relpath(target, out_root), processed_at=stamp)
            try:
                warns = write_spectrum(s, target, f)
                e.output_sha256 = sha256_file(target)
                e.output_size = target.stat().st_size
                if verify:
                    back = f.reader(target)
                    if len(back) < 1:
                        raise ValueError("output contains no spectrum")
                    v = verify_output(s, back[0], energy_tol=energy_tols[0], energy_fail=energy_tols[1],
                                     time_tol=time_tol)
                    e.verification = v
                    e.status = v.status
                    e.details = "; ".join(v.problems + v.warnings + [w for w in warns])
                else:
                    e.status = "WRITTEN"
                    e.details = "; ".join(warns)
            except Exception as exc:
                e.status = "ERROR"
                e.details = f"{type(exc).__name__}: {exc}"
            out.append(e)
    return out


def write_rescue_report(res: RescueResult, out: Path, manifest_rows: Optional[dict] = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    rows = []
    for e in res.entries:
        r = e.to_row()
        rows.append([str(e.source), r["spectrum_index"], r["format"], r["output_rel_path"], r["status"],
                     "skipped (verified earlier)" if e.skipped else "this run", r["counts_equal"],
                     _num(r["live_dev_s"]), _num(r["real_dev_s"]), _num(r["start_dev_s"]),
                     _num(r["energy_dev_keV"]), r["details"], r["source_sha256"], r["output_sha256"]])
    ws = _sheet(wb, "Verification", ["Source", "Spectrum #", "Format", "Output (relative)", "Status", "Run",
                                     "Counts identical", "Live-time dev (s)", "Real-time dev (s)",
                                     "Start-time dev (s)", "Max energy dev (keV)", "Details",
                                     "Source SHA-256", "Output SHA-256"], rows, first=True,
                widths={"Source": 60, "Output (relative)": 50, "Details": 60, "Source SHA-256": 20,
                        "Output SHA-256": 20, "Status": 20})
    fills = {"FAILED": "F8D7DA", "ERROR": "F8D7DA", "VERIFIED (warnings)": "FFF3CD", "VERIFIED": "D4EDDA"}
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, 5)
        if c.value in fills:
            c.fill = PatternFill("solid", fgColor=fills[c.value])
    s = wb.create_sheet("Summary")
    bold = Font(bold=True)
    for k, v in [("Archive root", str(res.root)), ("Output root", str(res.out_root)),
                 ("Source files considered", res.files_total), ("Elapsed (s)", round(res.elapsed, 1)),
                 ("Manifest", str(res.manifest)), ("Cancelled", "yes" if res.cancelled else "no")]:
        s.append([k, v])
    s.append([])
    s.append(["Status", "Outputs"])
    for k, v in sorted(res.counts().items()):
        s.append([k, v])
    per_fmt: dict[str, Counter] = defaultdict(Counter)
    edev: dict[str, float] = defaultdict(float)
    for e in res.entries:
        per_fmt[e.format_key][e.status] += 1
        try:
            edev[e.format_key] = max(edev[e.format_key], float(e.to_row()["energy_dev_keV"] or 0))
        except ValueError:
            pass
    s.append([])
    s.append(["Format", "Verified", "Verified with warnings", "Failed", "Error", "Max energy dev (keV)"])
    for k in sorted(per_fmt):
        c = per_fmt[k]
        s.append([k or "(none)", c["VERIFIED"], c["VERIFIED (warnings)"], c["FAILED"], c["ERROR"], edev[k]])
    for row in s.iter_rows():
        if row and row[0].value in ("Status", "Format") or (row and row[0].row <= 6):
            row[0].font = bold
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 60
    probs = [r for r in rows if r[4] in ("FAILED", "ERROR")]
    _sheet(wb, "Problems", ["Source", "Spectrum #", "Format", "Output (relative)", "Status", "Run",
                            "Counts identical", "Live-time dev (s)", "Real-time dev (s)", "Start-time dev (s)",
                            "Max energy dev (keV)", "Details", "Source SHA-256", "Output SHA-256"], probs,
           widths={"Source": 60, "Details": 80})
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _num(x):
    try:
        return float(x) if x not in ("", None) else None
    except ValueError:
        return x
