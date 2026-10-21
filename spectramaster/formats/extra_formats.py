"""Formats DataMaster never had: Excel workbook (code 50) and JSON (code 51).

Excel (.xlsx, needs openpyxl)
    Sheet "Spectrum": columns Channel, [Energy (keV)], Counts (energy column only if
    the spectrum is calibrated).  Sheet "Info": two columns (Field, Value) holding the
    metadata.  Files written by SpectraMaster can be read back; for other workbooks the
    reader looks for a header row containing "Counts" on the first sheet.

JSON (.json) -- schema version 1::

    {
      "format": "spectramaster-spectrum",
      "version": 1,
      "spectra": [
        {
          "sample_id": "...", "description": "...", "detector": "...",
          "start_time": "2016-07-27T14:06:47" | null,    # ISO 8601, naive local time
          "live_time": 450.0 | null, "real_time": 461.16 | null,   # seconds
          "first_channel": 0,
          "energy_coeffs": [c0, c1, c2],   # keV, E(ch) = c0 + c1*ch + c2*ch^2 ...
          "fwhm_coeffs": [],
          "counts": [0, 0, 12, ...],
          "extra": {"key": "value", ...}
        }
      ]
    }

    The reader also accepts a bare spectrum object or a list of them.
"""
from __future__ import annotations

import io
import json
import math
from datetime import datetime
from pathlib import Path
from typing import Any

from ..model import Spectrum
from .base import FormatError, FormatSpec, register

JSON_FORMAT_ID = "spectramaster-spectrum"
JSON_VERSION = 1


# ------------------------------------------------------------------ helpers
def _num(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer() and abs(v) < 2**53:
        return int(v)
    return v


def _jsonable(v: Any):
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, (list, tuple)):
        return [_jsonable(x) for x in v]
    if isinstance(v, dict):
        return {str(k): _jsonable(x) for k, x in v.items()}
    return str(v)


def _parse_time(v) -> datetime | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    try:
        return datetime.fromisoformat(str(v).strip()).replace(tzinfo=None)
    except ValueError:
        return None


def _float_or_none(v):
    if v in (None, ""):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _float_list(v) -> list[float]:
    if v in (None, ""):
        return []
    if isinstance(v, str):
        v = [p for p in v.replace(";", ",").replace("[", "").replace("]", "").split(",") if p.strip()]
    return [float(x) for x in v]


# --------------------------------------------------------------------- JSON
def spectrum_to_dict(s: Spectrum) -> dict:
    return {
        "sample_id": s.sample_id,
        "description": s.description,
        "detector": s.detector,
        "start_time": s.start_time.isoformat() if s.start_time else None,
        "live_time": s.live_time,
        "real_time": s.real_time,
        "first_channel": int(s.first_channel or 0),
        "energy_coeffs": [float(c) for c in s.energy_coeffs],
        "fwhm_coeffs": [float(c) for c in s.fwhm_coeffs],
        "counts": [_num(float(c)) for c in s.counts],
        "extra": _jsonable(dict(s.extra)),
    }


def spectrum_from_dict(d: dict) -> Spectrum:
    if not isinstance(d, dict) or "counts" not in d:
        raise FormatError("JSON spectrum object has no 'counts'")
    try:
        counts = [float(c) for c in d["counts"]]
    except (TypeError, ValueError):
        raise FormatError("'counts' must be a list of numbers") from None
    extra = d.get("extra") or {}
    return Spectrum(
        counts=counts,
        energy_coeffs=_float_list(d.get("energy_coeffs")),
        live_time=_float_or_none(d.get("live_time")),
        real_time=_float_or_none(d.get("real_time")),
        start_time=_parse_time(d.get("start_time")),
        sample_id=str(d.get("sample_id") or ""),
        description=str(d.get("description") or ""),
        detector=str(d.get("detector") or ""),
        first_channel=int(d.get("first_channel") or 0),
        fwhm_coeffs=_float_list(d.get("fwhm_coeffs")),
        extra=dict(extra) if isinstance(extra, dict) else {},
    )


def _json_write(s: Spectrum, path: Path) -> None:
    doc = {"format": JSON_FORMAT_ID, "version": JSON_VERSION, "spectra": [spectrum_to_dict(s)]}
    Path(path).write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding="utf-8")


def _json_read(path: Path) -> list[Spectrum]:
    try:
        doc = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise FormatError(f"{Path(path).name}: invalid JSON ({exc})") from None
    if isinstance(doc, dict) and "spectra" in doc:
        items = doc["spectra"]
    elif isinstance(doc, dict):
        items = [doc]
    elif isinstance(doc, list):
        items = doc
    else:
        raise FormatError(f"{Path(path).name}: unexpected JSON structure")
    return [spectrum_from_dict(d) for d in items]


def _json_sniff(path: Path, head: bytes) -> bool:
    h = head.lstrip(b"\xef\xbb\xbf \t\r\n")
    if not h.startswith((b"{", b"[")):
        return False
    return JSON_FORMAT_ID.encode() in head or (b'"counts"' in head and b'"energy_coeffs"' in head)


# -------------------------------------------------------------------- Excel
_INFO_FIELDS = [
    ("Sample ID", "sample_id"),
    ("Description", "description"),
    ("Detector", "detector"),
    ("Start time", "start_time"),
    ("Live time (s)", "live_time"),
    ("Real time (s)", "real_time"),
    ("Dead time (%)", None),
    ("Channels", None),
    ("First channel", "first_channel"),
    ("Total counts", None),
    ("Energy coefficients (keV)", "energy_coeffs"),
    ("FWHM coefficients (keV)", "fwhm_coeffs"),
    ("Source file", None),
    ("Source format", None),
]


def _xlsx_write(s: Spectrum, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:  # pragma: no cover
        raise FormatError("Excel output needs the 'openpyxl' package") from None
    wb = Workbook()
    ws = wb.active
    ws.title = "Spectrum"
    energies = s.energies()
    header = ["Channel"] + (["Energy (keV)"] if energies else []) + ["Counts"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    first = int(s.first_channel or 0)
    for i, c in enumerate(s.counts):
        c = _num(float(c))
        if energies:
            ws.append([first + i, round(energies[i], 6), c])
        else:
            ws.append([first + i, c])
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 12

    info = wb.create_sheet("Info")
    info.append(["Field", "Value"])
    for cell in info[1]:
        cell.font = Font(bold=True)
    values = {
        "Sample ID": s.sample_id,
        "Description": s.description,
        "Detector": s.detector,
        "Start time": s.start_time.isoformat(sep=" ") if s.start_time else "",
        "Live time (s)": s.live_time if s.live_time is not None else "",
        "Real time (s)": s.real_time if s.real_time is not None else "",
        "Dead time (%)": round(s.dead_time_percent, 4) if s.dead_time_percent is not None else "",
        "Channels": s.num_channels,
        "First channel": first,
        "Total counts": _num(s.total_counts),
        "Energy coefficients (keV)": ", ".join(repr(float(c)) for c in s.energy_coeffs),
        "FWHM coefficients (keV)": ", ".join(repr(float(c)) for c in s.fwhm_coeffs),
        "Source file": s.source_path,
        "Source format": s.source_format,
    }
    for label, _ in _INFO_FIELDS:
        info.append([label, values[label]])
    for k, v in s.extra.items():
        v = _jsonable(v)
        if isinstance(v, (list, dict)):
            v = json.dumps(v, ensure_ascii=False)
        info.append([f"extra.{k}", v])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 60
    wb.save(str(path))


def _xlsx_read(path: Path) -> list[Spectrum]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover
        raise FormatError("reading Excel needs the 'openpyxl' package") from None
    try:
        # from the bytes, not the name: openpyxl refuses names not ending in .xlsx/.xlsm
        wb = load_workbook(io.BytesIO(Path(path).read_bytes()), read_only=True, data_only=True)
    except Exception as exc:
        raise FormatError(f"{Path(path).name}: not a readable .xlsx workbook ({exc})") from None
    try:
        ws = wb["Spectrum"] if "Spectrum" in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        col_counts = col_chan = None
        for row in rows:
            labels = [str(v).strip().lower() if v is not None else "" for v in row]
            if "counts" in labels:
                col_counts = labels.index("counts")
                col_chan = labels.index("channel") if "channel" in labels else None
                break
        if col_counts is None:
            raise FormatError(f"{Path(path).name}: no 'Counts' column found")
        counts: list[float] = []
        first = None
        for row in rows:
            if row is None or len(row) <= col_counts or row[col_counts] is None:
                continue
            try:
                counts.append(float(row[col_counts]))
            except (TypeError, ValueError):
                continue
            if first is None and col_chan is not None and row[col_chan] is not None:
                try:
                    first = int(row[col_chan])
                except (TypeError, ValueError):
                    first = 0
        if not counts:
            raise FormatError(f"{Path(path).name}: the Counts column is empty")
        spec = Spectrum(counts=counts, first_channel=first or 0)
        if "Info" in wb.sheetnames:
            info = {}
            for row in wb["Info"].iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    info[str(row[0]).strip()] = row[1] if len(row) > 1 else None
            spec.sample_id = str(info.get("Sample ID") or "")
            spec.description = str(info.get("Description") or "")
            spec.detector = str(info.get("Detector") or "")
            spec.start_time = _parse_time(info.get("Start time"))
            spec.live_time = _float_or_none(info.get("Live time (s)"))
            spec.real_time = _float_or_none(info.get("Real time (s)"))
            try:
                spec.energy_coeffs = _float_list(info.get("Energy coefficients (keV)"))
                spec.fwhm_coeffs = _float_list(info.get("FWHM coefficients (keV)"))
            except ValueError:
                pass
            if info.get("First channel") not in (None, ""):
                try:
                    spec.first_channel = int(info["First channel"])
                except (TypeError, ValueError):
                    pass
            for k, v in info.items():
                if k.startswith("extra.") and v is not None:
                    spec.extra[k[6:]] = v
        return [spec]
    finally:
        wb.close()


def _xlsx_sniff(path: Path, head: bytes) -> bool:
    """A workbook written by us: zip containing a sheet called "Spectrum"."""
    if not head.startswith(b"PK\x03\x04"):
        return False
    import zipfile
    try:
        with zipfile.ZipFile(path) as zf:
            wb = zf.read("xl/workbook.xml")
    except Exception:
        return False
    return b'name="Spectrum"' in wb


register(FormatSpec(
    code=50, key="excel_xlsx", name="Microsoft Excel workbook", extensions=[".xlsx"],
    reader=_xlsx_read, writer=_xlsx_write, sniffer=_xlsx_sniff, aliases=["xlsx", "excel"],
    notes='Sheets "Spectrum" (Channel, Energy, Counts) and "Info" (metadata).'))

register(FormatSpec(
    code=51, key="json", name="SpectraMaster JSON", extensions=[".json"],
    reader=_json_read, writer=_json_write, sniffer=_json_sniff, aliases=["spectramaster_json"],
    notes="Documented schema mirroring the Spectrum data model (see extra_formats.py)."))
